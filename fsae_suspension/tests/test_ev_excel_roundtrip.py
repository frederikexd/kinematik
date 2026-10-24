"""The lap-sim ↔ EV workbook round trip, on a workbook built to the stock layout.

The stock FSAE_EV_Power_Draw.xlsx is not in the repository, so `_stock()` builds one
with the same structure: a SpeedVsTime profile with summary rows under it, an
ElecPropulsion sheet whose fifteen gear columns are live formulas in three stacked
blocks, ThermalLoad / EMFs / BearingBlowOut sheets that read those blocks by absolute
row (one through a Microsoft-365-only CHOOSECOLS), and a pack sheet read by label.

The checks are closed forms: the road load the lap asks for (pack current, energy,
motor speed and torque in the gear ThermalLoad selects), the workbook's own chain where
a flattened file is rebuilt, the layout contract (profile rows, block rows, summary rows
left alone), and the read-back paths' handling of the ways real files go wrong.
"""
import datetime as dt
import io
import math
import sys

import numpy as np
import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402

from suspension import ev_excel_roundtrip as rt  # noqa: E402

N = 40                                     # profile rows: SpeedVsTime 2..41
MPH = 2.23694
K_RPM = 1056.0 / (math.pi * 18.0)          # wheel 18 in: mph -> wheel rpm
M, CDA, CRR, RHO, ETA = 230.0, 1.10, 0.018, 1.225, 0.90
R_WHEEL = 18 * 0.0254 / 2


def _road(v, a=0.0, m=M):
    """F = m a + rho CdA v^2 / 2 + Crr m g, and the pack power max(F v, 0) / eta."""
    F = m * a + 0.5 * RHO * CDA * np.asarray(v, float) ** 2 + CRR * m * 9.81
    return F, np.maximum(F * v, 0.0) / ETA

EP_PARAMS = [("Motor Peak Torque (Nm)", 140), ("Motor Peak Power (kW)", 80),
             ("Switching Frequency (kHz)", 10), ("Poles", 10), ("Max DC (V)", 600),
             ("Motor Efficiency", 0.9545), ("Current From Pack (A)", 175),
             ("Pack Voltage (V)", 504), ("Max RPM", 6500), ("Wheel Diameter (in)", 18),
             ("Power Factor", 0.95)]

PACK_ROWS = [("Fuse Max (A)", 50), ("Parrallel Battery Count", 4),
             ("Series Battery Count", 140), ("Nominal Battery Voltage (V)", 3.6),
             ("Capacity Battery Cell (Ah)", 4.2), ("Endurance Length (km)", 22),
             ("Max Battery Cells", 600), ("Internal Resistance Battery Cell (ohm)", 0.012),
             ("Battery Cell Weight (kg)", 0.045), ("Battery Pack Cell Count", 560),
             ("Battery Pack Internal Resistance (ohm)", 1.68),      # the sheet's 4x error
             ("Battery Pack Nominal Voltage (V)", 504), ("Current Across a Single Cell (A)", 12.5),
             ("Power Draw (kW)", 80), ("Pack Capacity (Ah)", 16.8),
             ("Pack Energy (kWh)", 8.467), ("Joule Heating (kWh)", 0.85)]


def _stock(*, live=True, dependents=True, pack=True, svt_first=2, n=N, svt_name="SpeedVsTime",
           pack_name="Battery Pack Calcs", numeric_ratios=False):
    """A workbook in the stock layout, carrying the defects power_draw.audit_workbook pins.

    With n=1893, numeric_ratios=True and pack_name="BatteryPackConfig" the cell addresses
    are the stock file's, so the audit itself can be run on it.
    """
    wb = openpyxl.Workbook()
    svt = wb.active
    svt.title = svt_name
    if svt_first > 2:
        svt.cell(row=1, column=1, value="Lap profile")
    svt.cell(row=svt_first - 1, column=1, value="Time (s)")
    svt.cell(row=svt_first - 1, column=2, value="Speed (mph)")
    for i in range(n):
        svt.cell(row=svt_first + i, column=1, value=0.5 * i)
        svt.cell(row=svt_first + i, column=2, value=30.0 + i)
    last = svt_first + n - 1
    svt.cell(row=last + 2, column=1, value="Max Speed")
    svt.cell(row=last + 2, column=2, value=f"=MAX(B{svt_first}:B{last})")

    ep = wb.create_sheet("ElecPropulsion")
    for r, (lab, val) in enumerate(EP_PARAMS, start=1):
        ep.cell(row=r, column=1, value=lab)
        ep.cell(row=r, column=2, value=val)
    ep.cell(row=1, column=4, value="No Load Speed")
    ep.cell(row=1, column=5, value=7000)
    ep.cell(row=2, column=5, value=6000)
    ep.cell(row=1, column=8, value=1)                       # gear 1 is a literal 1:1
    for g in range(2, 16):                                  # the rest are 1/g
        ep.cell(row=1, column=7 + g, value=1 / g if numeric_ratios else f"=1/{g}")
    b2, b3 = 2 + n, 2 * n + 3
    for i in range(n):
        for c in range(8, 23):
            L = openpyxl.utils.get_column_letter(c)
            if live:
                ep.cell(row=2 + i, column=c,
                        value=f"=(SpeedVsTime!B{svt_first + i}*${L}$1*1056)/($B$10*PI())")
                ep.cell(row=b2 + i, column=c, value=f"=($B$8*$B$11*{L}{2 + i})/1000")
                ep.cell(row=b3 + i, column=c,
                        value=f"=($B$11*SQRT(3)*$B$6*{L}{b2 + i}*$B$8)/1000")
            else:
                ep.cell(row=2 + i, column=c, value=123.0)   # an earlier export's values
    if live:                                                # the summary rows, as shipped
        for si, (lab, r1, r2, agg) in enumerate([
                ("Min RPM", 2, n + 2, "MIN"), ("Max RPM", 2, n + 2, "MAX"),   # one row long
                ("Min Current (A)", b2, b2 + n - 1, "MIN"), ("Max Current (A)", b2, b2 + n - 1, "MAX"),
                ("Min Power (W)", b3, b3 + n - 1, "MIN"), ("Max Power (kW)", b3, b3 + n - 1, "MAX")]):
            r = b3 + n + 1 + si
            ep.cell(row=r, column=7, value=lab)
            for c in range(8, 23):
                L = openpyxl.utils.get_column_letter(c)
                ep.cell(row=r, column=c, value=f"={agg}({L}{r1}:{L}{r2})")

    q = pack_name if " " not in pack_name else f"'{pack_name}'"
    if dependents:
        tl = wb.create_sheet("ThermalLoad")
        tl["A2"], tl["B2"] = "Gear", "=1/7"
        tl["A4"], tl["B4"] = "Min Joule Heat", f"=MIN(F1:F{n})*{q}!B11"
        tl["A5"], tl["B5"] = "Max Joule Heat", f"=MAX(F1:F{n})*{q}!B11*{q}!B1"     # the stray fuse
        tl["A6"], tl["B6"] = "Sum of Joule Heat (kW)", f"=SUM(F1:F{n})"
        tl["E1"] = "Joule Heat (kW)"
        for r in range(1, n + 1):
            tl.cell(row=r, column=6,
                    value=f"=INDEX(ElecPropulsion!$H${b2}:$V${b2 + n - 1},{r},ROUND(1/$B$2,0))^2"
                          f"*{q}!B11")
        tl["D5"] = f"=MAX(ElecPropulsion!H{b2}:V{b2 + n - 1})"
        tl["D6"] = f"=_xlfn.CHOOSECOLS(ElecPropulsion!H2:V{n + 1},ROUND(1/B2,0))"
        tl["D7"] = ArrayFormula("D7:D9", f"=_xlfn.CHOOSEROWS(ElecPropulsion!H{b3}:V{b3 + n - 1},2)")
        emf = wb.create_sheet("EMFs")
        emf["A1"] = f"=SUM(ElecPropulsion!H{b3}:V{b3 + n - 1})"
        emf["B1"] = f"={q}!B12*ElecPropulsion!H{b3}:H{b3 + n}*{q}!B11"      # one row long
        bb = wb.create_sheet("BearingBlowOut")
        bb["A1"] = f"=MIN(ElecPropulsion!H2:V{n + 1})"
        bb["A16"], bb["B16"], bb["A17"], bb["B17"] = "Pitch diameter (m)", 0.05, "Pressure angle", 20
        for r in range(1, n + 1):
            bb.cell(row=r, column=7,
                    value=f"=(ElecPropulsion!N{b3 + r - 1}*9549/MAX(ElecPropulsion!N{r + 1},1)"
                          f"/$B$16)*TAN($B$17)")
        bb["A19"], bb["B19"] = "Max bearing load", f"=MAX(G1:G{n + 1})"
    if pack:
        bp = wb.create_sheet(pack_name)
        for r, (lab, val) in enumerate(PACK_ROWS, start=1):
            bp.cell(row=r, column=1, value=lab)
            bp.cell(row=r, column=2, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _load(b, **kw):
    return openpyxl.load_workbook(io.BytesIO(b), **kw)


def _const_lap(v=20.0, n=N, dt_s=0.5):
    return np.full(n, v), np.arange(n) * dt_s


# ------------------------------------------------------------------ sheet lookup
def test_sheets_are_found_by_alias_case_spacing_and_content():
    wb = openpyxl.Workbook()
    wb.active.title = "speed vs time"
    wb.create_sheet("Renamed")["A1"] = "Series Battery Count"
    wb["Renamed"]["A2"] = "Fuse Max"
    assert rt.resolve_svt_sheet(wb).title == "speed vs time"
    assert rt.resolve_pack_sheet(wb).title == "Renamed"
    assert rt.resolve_ep_sheet(wb) is None
    assert rt.resolve_sheet(wb, ("Nope",)) is None


def test_content_match_streams_when_random_access_is_refused():
    class Streaming:
        max_row = 3

        def cell(self, r, c):
            raise TypeError("read-only")

        def iter_rows(self, max_col, values_only):
            yield from [("Motor Peak Torque (Nm)",), ()] + [("x",)] * 50 + [("never read",)]
    wb = type("WB", (), {"sheetnames": ["s"], "worksheets": [Streaming()]})()
    assert rt.resolve_sheet(wb, ("ElecPropulsion",), ("motor peak torque (nm)",)) \
        is wb.worksheets[0]
    assert rt.resolve_sheet(wb, ("ElecPropulsion",), ("never read",)) is None

    class Short(Streaming):
        def iter_rows(self, max_col, values_only):
            yield ("fuse max",)
    wb2 = type("WB", (), {"sheetnames": ["s"], "worksheets": [Short()]})()
    assert rt.resolve_sheet(wb2, (), ("fuse max",)) is wb2.worksheets[0]


# ------------------------------------------------------------------- pack values
def test_pack_values_are_read_by_label_and_derived():
    warnings = []
    pack = rt.read_pack_values(_load(_stock(), data_only=True)["Battery Pack Calcs"], warnings)
    assert pack["fuse_max_a"] == 50 and pack["n_series"] == 140 and pack["n_parallel"] == 4
    assert pack["pack_voltage_v"] == pytest.approx(140 * 3.6)
    assert pack["pack_capacity_ah"] == pytest.approx(4 * 4.2)
    assert pack["pack_energy_wh"] == pytest.approx(504 * 16.8)
    assert pack["pack_energy_kwh"] == pytest.approx(504 * 16.8 / 1000)
    assert pack["pack_r_ohm"] == pytest.approx(140 * 0.012 / 4)        # S * R/P
    assert any("pack_r_ohm reads 1.68" in w for w in warnings)        # the sheet's error


def test_pack_values_flag_implausible_numbers_and_fall_back_to_coordinates():
    wb = openpyxl.Workbook()
    ws = wb.active
    for lab, val in [("Series Battery Count", 2000), ("Nominal Battery Voltage", 3.6), (None, 1),
                     ("Fuse Max", None)]:
        ws.append([lab, val])
    w = []
    pack = rt.read_pack_values(ws, w)
    assert "n_series" not in pack and "pack_voltage_v" not in pack
    assert sum("outside the plausible range" in x for x in w) == 2
    bare = openpyxl.Workbook().active
    for r in range(1, 17):
        bare.cell(row=r, column=2, value=float(r))
    w2 = []
    pk = rt.read_pack_values(bare, w2)
    assert pk["fuse_max_a"] == 1.0 and "no recognisable row labels" in w2[0]
    only_fuse = openpyxl.Workbook().active
    only_fuse.append(["Fuse Max", 60])
    assert rt.read_pack_values(only_fuse) == {"fuse_max_a": 60.0}       # nothing to derive


# ----------------------------------------------------------------- gear ratios
def test_gear_ratio_values():
    assert rt.gear_ratio_value(3) == (3.0, "ok")
    assert rt.gear_ratio_value(" 0.5 ") == (0.5, "ok")
    assert rt.gear_ratio_value(dt.time(3, 0)) == (0.125, "recovered")
    assert rt.gear_ratio_value(dt.timedelta(hours=2, minutes=24)) == (pytest.approx(0.1), "recovered")
    for bad in (None, True, dt.datetime(2026, 1, 1), dt.time(0, 0), "1/7?"):
        assert rt.gear_ratio_value(bad) == (None, "unreadable")


def test_gear_ratios_fall_back_to_the_formula_then_to_a_reported_default():
    data = _load(_stock(), data_only=True)["ElecPropulsion"]
    form = _load(_stock())["ElecPropulsion"]
    w = []
    ratios = rt.read_gear_ratios(data, form, w)
    assert ratios == pytest.approx([1.0] + [1 / g for g in range(2, 16)])
    assert "read from the cell formula" in w[0] and len(w) == 1
    form.cell(row=1, column=9).value = 0.25                      # numeric in the formula copy
    form.cell(row=1, column=10).value = "=A1*2"                  # not a simple reciprocal
    form.cell(row=1, column=11).value = None                     # nothing at all
    w = []
    ratios = rt.read_gear_ratios(data, form, w)
    assert ratios[1] == 0.25 and ratios[2] == 1.0 and ratios[3] == 1.0
    assert "3, 4 could not be read and were treated as 1:1" in w[1]
    assert rt.read_gear_ratios(None, None) == [1.0] * 15


def test_extracted_parameters_recover_formula_only_gear_ratios():
    """A workbook saved by openpyxl has no cached `=1/7`; the stored parameters must not read 1:1."""
    p = rt.extract_params_from_excel(_stock())
    assert p["gear_ratios"] == pytest.approx([1.0] + [1 / g for g in range(2, 16)])
    assert p["gear_ratios_recovered"] == list(range(2, 16))
    assert "read 0.142857 from its formula =1/7" in p["gear_ratio_notes"][5]
    assert p["motor"]["pack_voltage_ep_v"] == 504 and p["motor"]["synchronous_rpm"] == 6000
    assert p["pack"]["pack_voltage_v"] == pytest.approx(504)


def test_extracted_gear_ratios_report_dates_zeros_and_junk():
    wb = _load(_stock())
    ep = wb["ElecPropulsion"]
    ep.cell(row=1, column=9).value = dt.time(3, 0)            # a date-formatted 1/8
    ep.cell(row=1, column=10).value = 0                       # not a ratio
    ep.cell(row=1, column=11).value = "seven"
    buf = io.BytesIO()
    wb.save(buf)
    p = rt.extract_params_from_excel(buf.getvalue())
    assert p["gear_ratios"][1:4] == pytest.approx([0.125, 1.0, 1.0])
    assert 2 in p["gear_ratios_recovered"] and p["gear_ratios_unreadable"] == [3, 4]
    assert "date number format" in p["gear_ratio_notes"][0]
    assert "is 0, not a usable ratio" in p["gear_ratio_notes"][1]
    assert "could not be read ('seven')" in p["gear_ratio_notes"][2]


def test_extracted_parameters_without_sheets_or_openpyxl(monkeypatch):
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    p = rt.extract_params_from_excel(buf.getvalue())
    assert p["pack"] == {} and p["motor"] == {} and p["gear_ratios"] == [1.0] * 15
    assert "no propulsion sheet" in p["gear_ratio_notes"][0]
    assert "_error" in rt.extract_params_from_excel(b"not a workbook")
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert rt.extract_params_from_excel(b"")["_error"] == "openpyxl not installed"


# ------------------------------------------------------- profile extent & layout
def test_profile_extent_comes_from_the_propulsion_formulas():
    wb = _load(_stock())
    assert rt.native_profile_extent(wb, wb["ElecPropulsion"], wb["SpeedVsTime"]) == (2, N + 1)
    assert rt.dependent_block_starts(wb) == [2, N + 2, 2 * N + 3]
    assert rt._contract_block_height(wb) == N


def test_profile_extent_from_the_dependent_sheets_keeps_the_profile_origin():
    wb = _load(_stock(live=False, svt_first=4))
    ws_ep, ws_svt = wb["ElecPropulsion"], wb["SpeedVsTime"]
    assert rt.native_profile_extent(wb, ws_ep, ws_svt) == (4, N + 3)   # SVT origin, EP height
    assert rt._svt_first_data_row(None) == 2
    blank = openpyxl.Workbook().active
    assert rt._svt_first_data_row(blank, default=9) == 9


def test_profile_extent_from_the_populated_rows_starts_where_the_data_does():
    """Two title rows above the data used to be overwritten: the scan always started at row 2."""
    wb = _load(_stock(live=False, dependents=False, svt_first=4))
    assert rt.native_profile_extent(wb, wb["ElecPropulsion"], wb["SpeedVsTime"]) == (4, N + 3)
    v, t = _const_lap(15.0)
    buf = io.BytesIO()
    wb.save(buf)
    out = _load(rt.lap_to_excel_roundtrip(v, t, buf.getvalue()).excel_bytes)["SpeedVsTime"]
    assert out["A1"].value == "Lap profile" and out["A3"].value == "Time (s)"
    assert out["B4"].value == pytest.approx(15 * MPH, abs=1e-4)


def test_contract_height_goes_with_the_majority_and_ties_to_the_shorter():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(ElecPropulsion!H2:V41)"                  # 40
    ws["A2"] = "=SUM(ElecPropulsion!H42:V81)"                 # 40
    ws["A3"] = "=SUM(ElecPropulsion!H83:V123)"                # 41 (reads one row long)
    ws["A4"] = "=ElecPropulsion!H2:H3"                        # a summary span, ignored
    assert rt._contract_block_height(wb) == 40
    ws["A5"] = "=SUM(ElecPropulsion!H200:V240)"               # 41 again: a tie
    assert rt._contract_block_height(wb) == 40
    wb.create_sheet("ElecPropulsion")["A1"] = "=SUM(ElecPropulsion!H2:V999)"   # its own: skipped
    assert rt._contract_block_height(wb) == 40
    assert rt._contract_block_height(openpyxl.Workbook()) == 0


def test_profile_extent_falls_back_to_the_populated_rows_then_the_stock_layout():
    wb = _load(_stock(live=False, dependents=False))
    assert rt.native_profile_extent(wb, wb["ElecPropulsion"], wb["SpeedVsTime"]) == (2, N + 1)
    empty = openpyxl.Workbook()
    assert rt.native_profile_extent(empty, None, empty.active) == rt._SVT_FALLBACK_EXTENT
    assert rt.native_profile_extent(empty, None, None) == rt._SVT_FALLBACK_EXTENT
    empty.active.append([0.0, 12.0])                          # one sample is not a profile
    assert rt.native_profile_extent(empty, None, empty.active) == rt._SVT_FALLBACK_EXTENT


def test_resample_profile():
    assert rt.resample_profile([1, 2, 3], 0).size == 0
    assert rt.resample_profile([], 3) == pytest.approx([0, 0, 0])
    assert rt.resample_profile([7], 3) == pytest.approx([7, 7, 7])
    assert rt.resample_profile([1, 2, 3], 3) == pytest.approx([1, 2, 3])
    assert rt.resample_profile([0, 10], 5) == pytest.approx([0, 2.5, 5, 7.5, 10])


def test_portable_formula_rewrites_365_functions_exactly():
    f, n = rt.portable_formula('=SUM(_xlfn.CHOOSECOLS(A1:C9,2))+_xlfn.CHOOSEROWS(B1:B9, 3)')
    assert (f, n) == ("=SUM(INDEX(A1:C9,0,2))+INDEX(B1:B9,3,0)", 2)
    f, n = rt.portable_formula('=_xlfn.CHOOSECOLS(IF(A1=",",B1:C2,D1:E2),MAX(1,2))')
    assert f == '=INDEX(IF(A1=",",B1:C2,D1:E2),0,MAX(1,2))' and n == 1
    assert rt.portable_formula("=_xlfn.CHOOSECOLS(A1:B2)") == ("=_xlfn.CHOOSECOLS(A1:B2)", 0)
    assert rt.portable_formula("=_xlfn.CHOOSECOLS(A1:B2,") == ("=_xlfn.CHOOSECOLS(A1:B2,", 0)
    assert rt.portable_formula("=SUM(A1)") == ("=SUM(A1)", 0)


def test_make_workbook_portable_rewrites_plain_and_array_formulas(monkeypatch):
    wb = _load(_stock())
    assert rt.make_workbook_portable(wb) == 2
    assert wb["ThermalLoad"]["D6"].value.startswith("=INDEX(ElecPropulsion!H2:V41,0,")
    arr = wb["ThermalLoad"]["D7"].value
    assert isinstance(arr, ArrayFormula) and arr.text.startswith("=INDEX(") and arr.ref == "D7:D9"
    wb["ThermalLoad"]["C1"] = "=_xlfn.XLOOKUP(1,A1:A2,B1:B2)"   # 365-only but not ours to change
    assert rt.make_workbook_portable(wb) == 0
    assert wb["ThermalLoad"]["C1"].value.startswith("=_xlfn.XLOOKUP")
    monkeypatch.setitem(sys.modules, "openpyxl.worksheet.formula", None)
    assert rt.make_workbook_portable(wb) == 0
    assert rt.dependent_block_ranges(wb)[0] == (2, N + 1)            # plain strings still read


def test_small_helpers():
    assert [rt._col_letter(n) for n in (1, 8, 22, 26, 27, 703)] == ["A", "H", "V", "Z", "AA", "AAA"]
    assert rt._safe_float("x", 5) == 5 and rt._safe_float(None) == 0.0 and rt._safe_float("2") == 2
    t = rt.lap_result_to_time_axis(np.array([10.0, 10.0, 0.0]), np.array([0.0, 50.0, 50.25]))
    assert t == pytest.approx([0, 5, 5 + 0.25 / 5])           # slow samples use a 0.5 m/s floor
    rt._request_full_recalc(object())                         # no calculation props: no crash


# --------------------------------------------------------- speed-profile loading
def test_speed_profile_loads_from_bytes_and_degrades_to_a_flat_line(tmp_path):
    t, v = rt.load_speed_vs_time_from_bytes(_stock())
    assert t == pytest.approx(0.5 * np.arange(N)) and v == pytest.approx(30.0 + np.arange(N))
    wb = openpyxl.Workbook()
    p = tmp_path / "no_svt.xlsx"
    wb.save(p)
    for path in (str(p), str(tmp_path / "missing.xlsx")):
        tt, vv = rt.load_speed_vs_time(path)
        assert tt == pytest.approx([0, 1]) and vv == pytest.approx([0, 0])


def test_speed_csv_shim_calls_the_electrical_check():
    from suspension.ev_electrical_check import ElecParams
    a = rt.check_lap_from_speed_csv([20.0] * 10, list(range(10)), ElecParams())
    from suspension.ev_electrical_check import check_lap_from_speed_csv
    b = check_lap_from_speed_csv([20.0] * 10, list(range(10)), ElecParams())
    assert (a.ok, a.peak_current_a) == (b.ok, b.peak_current_a)


# ---------------------------------------------------------------- the lap chain
def _closed_form(v_ms, pack_v=504.0, pf=0.95, ratio=1.0):
    """The workbook's own chain: RPM in a gear column and its speed-only current."""
    rpm = v_ms * MPH * ratio * K_RPM
    return rpm, pack_v * pf * rpm / 1000.0


def test_constant_speed_current_is_the_road_load_over_the_pack_voltage():
    v, t = _const_lap(20.0)
    r = rt.lap_to_excel_roundtrip(v, t, _stock(), lap_time_s=19.5, top_speed_ms=20.0,
                                  avg_speed_ms=20.0)
    assert r.ok, r.error
    F, P = _road(20.0)
    assert r.accel_ms2 == pytest.approx(np.zeros(N))
    assert r.current_draw_a == pytest.approx(np.full(N, P / 504.0))
    assert r.power_kw == pytest.approx(np.full(N, P / 1000.0))
    assert r.total_energy_kwh == pytest.approx(P * t[-1] / 3.6e6)         # P x lap duration
    assert r.usable_energy_kwh == pytest.approx(504 * 16.8 * 0.92 / 1000)
    # Gear 7 from ThermalLoad!B2 = 1/7: the motor turns seven times the wheel.
    assert (r.gear, r.gear_reduction) == (7, pytest.approx(7.0))
    assert r.wheel_rpm == pytest.approx(np.full(N, 20 * MPH * K_RPM))
    assert r.motor_rpm == pytest.approx(np.full(N, 7 * 20 * MPH * K_RPM))
    assert r.motor_torque_nm == pytest.approx(np.full(N, F * R_WHEEL / 7 / ETA))
    # The workbook's formula in the gear-7 column, for comparison only.
    assert r.workbook_current_a == pytest.approx(np.full(N, _closed_form(20.0, ratio=1 / 7)[1]))
    assert r.feasible and r.fuse_ok and r.motor_ok
    assert r.verdict.startswith("✅ Electrically feasible") and "gear 7" in r.verdict
    road = lambda u: (0.5 * RHO * CDA * u ** 3 + CRR * M * 9.81 * u)
    cap = 50 * 504 * ETA
    assert road(r.fuse_speed_ceiling_ms) >= cap > road(r.fuse_speed_ceiling_ms - 0.03)
    assert r.max_speed_mph == pytest.approx(20 * MPH)


def test_acceleration_adds_m_a_and_braking_draws_nothing():
    t = np.arange(0, 16.0, 0.1)
    v = np.where(t < 10, 5 + 2.0 * t, 25 - 3.0 * (t - 10))       # +2 m/s^2, then -3 m/s^2 to 7
    r = rt.lap_to_excel_roundtrip_from_db(v, t, rt.extract_params_from_excel(_stock()))
    accel = (t > 0.15) & (t < 9.85)
    F, P = _road(v[accel], 2.0)
    assert r.accel_ms2[accel] == pytest.approx(np.full(accel.sum(), 2.0))
    assert r.current_draw_a[accel] == pytest.approx(P / 504.0)
    assert np.all(r.current_draw_a[(t > 10.15)] == 0.0)          # 3 m/s^2 is past the road load
    assert r.motor_torque_nm[accel] == pytest.approx(F * R_WHEEL / 7 / ETA)
    heavy = rt.lap_to_excel_roundtrip_from_db(v, t, rt.extract_params_from_excel(_stock()),
                                              vehicle={"vehicle_mass_kg": 460.0})
    assert heavy.current_draw_a[accel] == pytest.approx(_road(v[accel], 2.0, m=460.0)[1] / 504.0)


def test_the_speed_only_formula_misses_the_launch():
    """Same speed, different load: the workbook formula reads both the same; the road load does not."""
    t = np.arange(0, 4.0, 0.1)
    launch = 10 + 3.0 * t
    cruise = np.full(t.size, 10 + 3.0 * 2.0)
    p = rt.extract_params_from_excel(_stock())
    a = rt.lap_to_excel_roundtrip_from_db(launch, t, p)
    b = rt.lap_to_excel_roundtrip_from_db(cruise, t, p)
    i = 20                                                       # both at 16 m/s here
    assert a.workbook_current_a[i] == pytest.approx(b.workbook_current_a[i])
    assert a.current_draw_a[i] > 2.5 * b.current_draw_a[i]


def test_motor_limits_are_checked_in_the_selected_gear():
    p = rt.extract_params_from_excel(_stock())
    fast = rt.lap_to_excel_roundtrip_from_db(*_const_lap(30.0), p)
    assert not fast.motor_ok and not fast.feasible
    assert f"motor over its max speed in gear 7 ({7 * 30 * MPH * K_RPM:.0f} vs 6500 rpm)" \
        in fast.verdict
    t = np.arange(0, 2.0, 0.05)
    hard = rt.lap_to_excel_roundtrip_from_db(2 + 9.0 * t, t, p, gear=1)    # 1:1, a hard launch
    assert "motor torque short in gear 1 (" in hard.verdict
    weak = dict(p, motor=dict(p["motor"], motor_peak_power_kw=5.0))
    assert "motor power short" in rt.lap_to_excel_roundtrip_from_db(*_const_lap(20.0), weak).verdict
    low = rt.lap_to_excel_roundtrip_from_db(*_const_lap(30.0), p, gear=3)
    assert low.motor_ok and low.gear == 3                          # 3:1 keeps it under 6500 rpm


def test_selected_gear_is_read_from_thermalload():
    wb_f, wb_d = _load(_stock()), _load(_stock(), data_only=True)
    assert rt.selected_gear(wb_d, wb_f) == 7                        # from the formula =1/7
    wb_f["ThermalLoad"]["B2"] = 0.2
    assert rt.selected_gear(None, wb_f) == 5                        # a cached number
    wb_f["ThermalLoad"]["B2"] = dt.time(2, 24)                      # 1/10, date-formatted
    assert rt.selected_gear(wb_f) == 10
    for bad in ("=1/40", 0, "=A1", None):
        wb_f["ThermalLoad"]["B2"] = bad
        assert rt.selected_gear(wb_f) is None
    assert rt.selected_gear(_load(_stock(dependents=False))) is None
    assert rt.extract_params_from_excel(_stock())["selected_gear"] == 7


def test_without_a_selector_gear_1_is_used_and_said():
    v, t = _const_lap(10.0)
    r = rt.lap_to_excel_roundtrip(v, t, _stock(dependents=False))
    assert r.gear == 1 and any("No gear selector found" in w for w in r.warnings)
    d = rt.lap_to_excel_roundtrip_from_db(v, t, {})
    assert d.gear == 1 and "No gear selector found" in d.warnings[0]
    flat = rt.lap_to_excel_roundtrip_from_db(v, np.zeros(N), {"selected_gear": 7})
    assert flat.accel_ms2 == pytest.approx(np.zeros(N))
    assert flat.warnings == ["The time axis is not strictly increasing, so the acceleration "
                             "could not be taken from it: pack current is the steady-speed "
                             "road load only, with no acceleration demand."]


def test_round_trip_from_the_database_agrees_with_the_workbook_path():
    v = 8.0 + 4.0 * np.sin(np.linspace(0, 3, N)) ** 2
    t = np.arange(N) * 0.5
    params = rt.extract_params_from_excel(_stock())
    a = rt.lap_to_excel_roundtrip(v, t, _stock())
    b = rt.lap_to_excel_roundtrip_from_db(v, t, params)
    for f in ("current_draw_a", "power_kw", "motor_rpm", "workbook_current_a", "accel_ms2"):
        assert getattr(b, f) == pytest.approx(getattr(a, f))
    assert (b.total_energy_kwh, b.verdict, b.excel_bytes) == (a.total_energy_kwh, a.verdict, b"")
    assert b.feasible and b.verdict.startswith("✅ Electrically feasible")


def test_energy_deficit_verdict_and_defaults():
    v, t = _const_lap(3.0, n=3000, dt_s=1.0)
    tiny = {"pack": {"pack_capacity_ah": 0.01, "fuse_max_a": 500.0}, "motor": {}}
    r = rt.lap_to_excel_roundtrip_from_db(v, t, tiny)
    assert r.fuse_ok and not r.energy_ok
    assert r.verdict.startswith("❌ NOT feasible — energy deficit")
    ramp_t = np.arange(0, 10.0, 0.1)
    both = rt.lap_to_excel_roundtrip_from_db(3.0 * ramp_t, ramp_t,
                                             {"pack": {"pack_capacity_ah": 0.001}})
    assert "fuse blown" in both.verdict and "energy deficit" in both.verdict
    short = rt.lap_to_excel_roundtrip_from_db([1.0], [0.0], {})
    assert not short.ok and "≥2" in short.error
    odd = rt.lap_to_excel_roundtrip_from_db([1.0, 2.0, 3.0], [0.0], {})
    assert odd.time_s == pytest.approx([0, 0.1, 0.2])


# ------------------------------------------------------------- the written file
def test_round_trip_writes_the_profile_and_leaves_the_model_alone():
    src = _stock()
    v = np.linspace(10, 30, 25)                       # 25 samples into 40 rows
    t = np.arange(25) * 0.8
    r = rt.lap_to_excel_roundtrip(v, t, src, lap_time_s=19.2, top_speed_ms=30, avg_speed_ms=20,
                                  correct_workbook=False)
    assert r.ok and any("Resampled to fit" in w for w in r.warnings) and r.corrections == []
    out = _load(r.excel_bytes)
    svt = out["SpeedVsTime"]
    speeds = [svt.cell(row=2 + i, column=2).value for i in range(N)]
    assert speeds == pytest.approx(rt.resample_profile(v * MPH, N), abs=1e-4)
    assert svt.cell(row=N + 3, column=2).value == f"=MAX(B2:B{N + 1})"          # user's row kept
    before = _load(src)["ElecPropulsion"]
    ep = out["ElecPropulsion"]
    assert all(ep.cell(row=rr, column=c).value == before.cell(row=rr, column=c).value
               for rr in range(1, 3 * N + 9) for c in range(8, 23))
    assert not any("_xlfn." in str(c.value) for row in out["ThermalLoad"].iter_rows() for c in row)
    assert any("Rewrote 2 Microsoft-365-only" in w for w in r.warnings)
    assert out.calculation.fullCalcOnLoad is True
    bp = out["Battery Pack Calcs"]
    start = len(PACK_ROWS) + 2
    block = {bp.cell(row=start + k, column=1).value: bp.cell(row=start + k, column=2).value
             for k in range(20)}
    assert block["Lap Time (s)"] == 19.2 and block["Top Speed (km/h)"] == 108.0
    assert block["Profile Points"] == 25 and block["Feasibility"] == ("PASS" if r.feasible else "FAIL")
    assert block["Peak Current Draw (A)"] == round(r.peak_current_a, 2)
    assert block["Gear (ThermalLoad selector)"] == "7 (7:1)"
    assert block["Max Motor Speed (rpm)"] == round(float(np.max(r.motor_rpm)), 0)
    assert block["Workbook Current Formula Peak (A)"] == round(float(np.max(r.workbook_current_a)), 2)


def test_a_flattened_propulsion_sheet_is_rebuilt_at_the_rows_its_readers_name():
    v, t = _const_lap(15.0)
    r = rt.lap_to_excel_roundtrip(v, t, _stock(live=False), correct_workbook=False)
    assert r.ok and any("Rebuilt all three blocks at rows 2-" in w for w in r.warnings)
    ep = _load(r.excel_bytes)["ElecPropulsion"]
    rpm, cur = _closed_form(15.0)
    ratios = [1.0] + [1 / g for g in range(2, 16)]
    b2, b3 = N + 2, 2 * N + 3                                     # the dependents' starts
    for gi, ratio in enumerate(ratios):
        c = 8 + gi
        assert ep.cell(row=2, column=c).value == pytest.approx(rpm * ratio, abs=1e-3)
        cur_g = 504 * 0.95 * rpm * ratio / 1000
        assert ep.cell(row=b2, column=c).value == pytest.approx(cur_g, abs=1e-5)
        pw_g = 0.95 * math.sqrt(3) * 0.9545 * cur_g * 504 / 1000
        assert ep.cell(row=b3 + N - 1, column=c).value == pytest.approx(pw_g, abs=1e-5)
    assert [ep.cell(row=r_, column=7).value for r_ in (2, b2, b3)] == \
        ["RPM (Load)", "Current Draw (A)", "Motor Power (kW)"]
    assert ep.cell(row=b3 + N + 1, column=7).value == "Min RPM"
    assert ep.cell(row=b3 + N + 2, column=14).value == f"=MAX(N2:N{N + 1})"     # gear 7 column


def test_rebuild_without_a_usable_contract_uses_the_stock_spacing():
    v, t = _const_lap(15.0)
    r = rt.lap_to_excel_roundtrip(v, t, _stock(live=False, dependents=False),
                                  correct_workbook=False)
    assert r.ok and not any("do not fit" in w for w in r.warnings)
    ep = _load(r.excel_bytes)["ElecPropulsion"]
    assert ep.cell(row=N + 2, column=7).value == "Current Draw (A)"
    assert ep.cell(row=2 * N + 3, column=7).value == "Motor Power (kW)"

    wb = _load(_stock(live=False, dependents=False))
    ws = wb.create_sheet("EMFs")
    ws["A1"] = f"=SUM(ElecPropulsion!H2:V{N + 1})"
    ws["A2"] = f"=SUM(ElecPropulsion!H10:V{N + 9})"               # overlaps the first block
    buf = io.BytesIO()
    wb.save(buf)
    r2 = rt.lap_to_excel_roundtrip(v, t, buf.getvalue(), correct_workbook=False)
    assert any("do not fit 40-row blocks" in w for w in r2.warnings)
    r3 = rt.lap_to_excel_roundtrip(v, t, buf.getvalue())
    assert any("do not fit 40-row blocks" in w for w in r3.warnings)


def test_round_trip_refuses_or_warns_on_missing_sheets(monkeypatch):
    v, t = _const_lap()
    no_svt = _load(_stock())
    no_svt["SpeedVsTime"].title = "Laps"
    no_svt["Laps"]["A1"] = "t"
    buf = io.BytesIO()
    no_svt.save(buf)
    r = rt.lap_to_excel_roundtrip(v, t, buf.getvalue())
    assert not r.ok and "no speed-profile sheet" in r.error
    no_ep = _load(_stock(dependents=False))
    del no_ep["ElecPropulsion"]
    buf = io.BytesIO()
    no_ep.save(buf)
    r = rt.lap_to_excel_roundtrip(v, t, buf.getvalue())
    assert not r.ok and "no propulsion sheet" in r.error
    r = rt.lap_to_excel_roundtrip(v, t, _stock(pack=False))
    assert r.ok and any("Pack sheet not found" in w for w in r.warnings)
    assert r.usable_energy_kwh == pytest.approx(504 * 15 * 0.92 / 1000)      # default pack
    assert not rt.lap_to_excel_roundtrip([1.0], [0.0], b"").ok
    odd = rt.lap_to_excel_roundtrip([5.0] * N, [0.0], _stock())
    assert odd.time_s == pytest.approx(np.arange(N) * 0.1)
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert rt.lap_to_excel_roundtrip(v, t, b"").error == "openpyxl not installed."


# ----------------------------------------------------------- extended analysis
def test_cell_thermals_heat_balance():
    t = np.arange(0, 600.0, 1.0)
    i_pack = np.full(t.size, 40.0)
    th = rt.compute_cell_thermals(i_pack, t, {"n_parallel": 4, "cell_r_ohm": 0.01,
                                             "cell_weight_kg": 0.05}, ambient_c=30.0)
    q = (40 / 4) ** 2 * 0.01                                         # 1 W per cell
    assert th["i_cell_a"] == pytest.approx(np.full(t.size, 10.0))
    assert th["peak_joule_w"] == pytest.approx(q)
    assert th["total_joule_kwh"] == pytest.approx(q * 599 / 3.6e6)
    C, hA = 0.05 * 900, 10 * 0.002
    T = 30.0
    for _ in range(599):                                              # explicit Euler, by hand
        T += (q - hA * (T - 30.0)) / C
    assert th["temp_c"][-1] == pytest.approx(T) and th["peak_temp_c"] == pytest.approx(T)
    assert T < 30 + q / hA                                            # below the steady state
    d = rt.compute_cell_thermals(i_pack[:2], t[:2], {}, thermal_mass_j_k=1.0)
    assert d["joule_w"][0] == pytest.approx(40 ** 2 * 0.003)          # defaults: 1P, 3 mOhm


def test_soc_finishes_or_stops_where_the_energy_runs_out():
    t = np.arange(0, 101.0)
    pack = {"pack_voltage_v": 400.0, "pack_capacity_ah": 10.0}
    usable = 400 * 10 * 0.92 / 1000                                   # 3.68 kWh
    p_ok = np.full(t.size, usable * 3600 / 200)                       # half the pack in 100 s
    s = rt.compute_soc_and_stop(p_ok, t, pack, {})
    assert s["finishes"] and s["stop_idx"] is None and s["deficit_kwh"] == 0.0
    assert s["soc_pct"][-1] == pytest.approx(50.0)
    p_bad = np.full(t.size, usable * 3600 / 49.5)                     # empty after 49.5 s
    s = rt.compute_soc_and_stop(p_bad, t, pack, {})
    assert not s["finishes"] and s["stop_idx"] == 50 and s["stop_time_s"] == 50.0
    assert s["pct_lap_done"] == pytest.approx(0.5)
    assert s["deficit_kwh"] == pytest.approx(usable * (100 / 49.5 - 1))
    assert s["soc_pct"][-1] == 0.0
    backwards = rt.compute_soc_and_stop(np.full(2, 1e9), np.array([5.0, 0.0]), pack, {})
    assert backwards["stop_idx"] == 1 and backwards["pct_lap_done"] == 0.0


def test_minimum_feasible_pack_sizes_to_usable_energy():
    t = np.arange(0, 1001.0)
    p = np.full(t.size, 18.0)                                         # 18 kW for 1000 s = 5 kWh
    i = np.full(t.size, 40.0)
    pack = {"pack_voltage_v": 504.0, "cell_voltage_v": 3.6, "cell_capacity_ah": 4.2,
            "cell_weight_kg": 0.045, "fuse_max_a": 50.0, "pack_capacity_ah": 16.8,
            "n_parallel": 4}
    m = rt.compute_minimum_feasible_pack(p, t, i, pack, {})
    assert m["min_energy_kwh"] == 5.0 and m["rec_energy_kwh"] == 5.5
    assert m["min_capacity_ah"] == pytest.approx(5000 / (504 * 0.92), abs=0.01)
    assert m["rec_cells_series"] == 140                               # derived from 504 V / 3.6 V
    assert m["min_cells_parallel"] == math.ceil(5000 / (504 * 0.92) / 4.2)
    assert m["rec_total_cells"] == 140 * m["rec_cells_parallel"]
    assert m["cur_cells_parallel"] == 4 and m["cur_total_cells"] == 560
    assert m["current_usable_kwh"] == pytest.approx(504 * 16.8 * 0.92 / 1000, abs=1e-4)
    assert m["energy_shortfall_kwh"] == 0.0 and m["fuse_ok"]
    # The minimum pack's usable energy is exactly what the lap draws.
    usable_min = m["min_capacity_ah"] * 504 * 0.92 / 1000
    assert usable_min == pytest.approx(m["min_energy_kwh"], rel=1e-3)
    small = rt.compute_minimum_feasible_pack(p, t, i * 2, {"pack_capacity_ah": 2.0}, {})
    assert small["energy_shortfall_kwh"] > 0 and not small["fuse_ok"]
    assert small["cur_cells_parallel"] == 1 and small["rec_cells_series"] == 140


# ------------------------------------------------------------ the enhanced file
@pytest.fixture(scope="module")
def enhanced():
    t = np.arange(335) * 0.04
    v = 10.0 + 5.0 * np.sin(np.linspace(0.0, 7.0, 335)) ** 2
    base = rt.lap_to_excel_roundtrip(v, t, _stock(), lap_time_s=float(t[-1]))
    th = rt.compute_cell_thermals(base.current_draw_a, base.time_s, base.pack)
    soc = rt.compute_soc_and_stop(base.power_kw, base.time_s, base.pack, base.motor)
    mp = rt.compute_minimum_feasible_pack(base.power_kw, base.time_s, base.current_draw_a,
                                          base.pack, base.motor)
    once = rt.build_enhanced_excel(base, base.excel_bytes, th, soc, mp)
    twice = rt.build_enhanced_excel(base, once, th, soc, mp)          # re-export replaces sheets
    return base, th, soc, mp, _load(twice)


def test_enhanced_heatmap_states_the_modelled_temperature(enhanced):
    base, th, _soc, _mp, wb = enhanced
    assert wb.sheetnames.count("PackHeatmap") == 1
    ws = wb["PackHeatmap"]
    grid = [ws.cell(row=7 + s, column=2 + p).value for s in range(140) for p in range(4)]
    assert set(grid) == {round(th["peak_temp_c"], 1)}                 # no invented gradient
    assert "one peak temperature for all of them" in ws.cell(row=7 + 140, column=1).value
    assert ws.cell(row=3, column=4).value == "140S × 4P"
    first = 7 + 140 + 2 + 3 + 1
    assert ws.cell(row=first, column=1).value == 0.0
    assert ws.cell(row=first, column=3).value == round(float(base.current_draw_a[0]), 3)
    assert "refresh_heatmap()" in [c.value for c in ws["A"]]


def test_enhanced_lap_energy_and_advisor(enhanced):
    base, _th, soc, mp, wb = enhanced
    le = wb["LapEnergy"]
    card = {le.cell(row=2, column=c).value: le.cell(row=3, column=c).value for c in range(1, 8)}
    assert card["Pack finishes lap"].endswith("YES") and "finishes" in card["Stop time"]
    assert card["Total energy drawn"] == f"{base.total_energy_kwh:.3f} kWh"
    assert len(le._charts) == 1
    fp = wb["FeasiblePack"]
    rows = {fp.cell(row=r, column=1).value: [fp.cell(row=r, column=c).value for c in (2, 3, 4)]
            for r in range(3, 12)}
    assert rows["Energy (usable)"] == [mp["current_usable_kwh"], mp["min_energy_kwh"],
                                       mp["rec_energy_kwh"]]
    assert rows["Energy (usable)"][1] == pytest.approx(base.total_energy_kwh, abs=1e-4)
    assert rows["Fuse OK for this lap"] == ["✅ OK"] * 3
    caps = [fp.cell(row=16 + k, column=1).value for k in range(9)]
    assert caps == [5, 8, 10, 12, 15, 18, 20, 25, 30]
    laps5 = 504 * 5 * 0.92 / 1000 / base.total_energy_kwh
    assert fp.cell(row=16, column=3).value == round(laps5, 2)


def test_enhanced_file_marks_the_stop_and_the_shortfall():
    t = np.arange(400) * 1.0
    v = np.full(400, 20.0)
    base = rt.lap_to_excel_roundtrip_from_db(v, t, {"pack": {"pack_capacity_ah": 0.5,
                                                             "fuse_max_a": 10.0}})
    th = rt.compute_cell_thermals(base.current_draw_a, t, base.pack)
    soc = rt.compute_soc_and_stop(base.power_kw, t, {"pack_capacity_ah": 0.5}, {})
    mp = rt.compute_minimum_feasible_pack(base.power_kw, t, base.current_draw_a,
                                          {"pack_capacity_ah": 0.5, "fuse_max_a": 10.0}, {})
    assert not soc["finishes"] and not mp["fuse_ok"]
    wb = _load(rt.build_enhanced_excel(base, _stock(), th, soc, mp))
    le = wb["LapEnergy"]
    markers = [le.cell(row=r, column=7).value for r in range(6, 6 + 200)]
    assert markers.count("🛑 STOP") == 1
    card = {le.cell(row=2, column=c).value: le.cell(row=3, column=c).value for c in range(1, 8)}
    assert card["Pack finishes lap"].endswith("NO") and card["Stop time"].endswith(" s")
    fp = wb["FeasiblePack"]
    rows = {fp.cell(row=r, column=1).value: fp.cell(row=r, column=2) for r in range(3, 12)}
    assert rows["Energy shortfall"].font.color.rgb == "FFE74C3C"
    assert rows["Fuse OK for this lap"].value == "❌ Over"


def test_enhanced_export_without_openpyxl_returns_the_input(monkeypatch):
    base = rt.lap_to_excel_roundtrip_from_db([1.0, 2.0], [0.0, 1.0], {})
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert rt.build_enhanced_excel(base, b"orig", {}, {}, {}) == b"orig"


# -------------------------------------------------- raw sheets & ledger handoff
def test_raw_sheets_echo_numeric_columns(monkeypatch):
    out = rt.read_raw_sheets(_stock())
    svt = [s for s in out["sheets"] if s["name"] == "SpeedVsTime"][0]
    assert [c["header"] for c in svt["columns"]] == ["Time (s)", "Speed (mph)"]
    assert svt["columns"][1]["values"][:3] == [30.0, 31.0, 32.0]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([1, "2", "1,000", True, dt.datetime(2026, 1, 1)])
    ws.append([3, "4", "x", None, dt.datetime(2026, 1, 2)])
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    raw = rt.read_raw_sheets(buf.getvalue())
    cols = raw["sheets"][0]["columns"]
    assert [c["header"] for c in cols] == ["Col A", "Col B"] and cols[1]["values"] == [2.0, 4.0]
    assert raw["sheet_names"] == ["Sheet", "Empty"] and len(raw["sheets"]) == 1
    assert rt._col_letter(1) == cols[0]["column_letter"]
    capped = rt.read_raw_sheets(_stock(), max_rows=3)
    svt3 = [s for s in capped["sheets"] if s["name"] == "SpeedVsTime"][0]
    assert svt3["n_rows"] == 2 and svt3["columns"][0]["values"] == [0.0, 0.5]
    sheet = rt.RawSheet(name="s", n_rows=1, n_cols=1,
                        columns=[rt.RawColumn("s", "h", "A", [1.0], 2)])
    assert sheet.column_headers() == ["h"]
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert rt.read_raw_sheets(b"")["_error"] == "openpyxl not installed"


def test_feasible_pack_sheet_is_surfaced_verbatim(monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Minimum Pack"
    ws.append(["Energy", 7.123456, None])
    ws.append([None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    f = rt.find_feasible_pack_sheet(buf.getvalue())
    assert f == {"found": True, "sheet": "Minimum Pack", "rows": [["Energy", 7.1235, ""]]}
    assert rt.find_feasible_pack_sheet(_stock()) == {"found": False, "sheet": "", "rows": []}
    for r in range(300):
        ws.append([r])
    buf = io.BytesIO()
    wb.save(buf)
    assert len(rt.find_feasible_pack_sheet(buf.getvalue())["rows"]) == 200
    assert "_error" in rt.find_feasible_pack_sheet(b"junk")
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert rt.find_feasible_pack_sheet(b"")["_error"] == "openpyxl not installed"


def test_ledger_declarations_from_the_stock_workbook():
    d = rt.ledger_declarations_from_ev_params(rt.extract_params_from_excel(_stock()))
    assert d["powertrain"] == {"peak_power_kw": 80, "peak_torque_nm": 140,
                               "voltage_v": pytest.approx(504)}
    assert d["electrics"]["peak_current_a"] == 175                # the pack's, not one cell's
    assert d["electrics"]["power_draw_w"] == 80000
    stint_s = 22 / 50 * 3600
    assert d["electrics"]["heat_reject_w"] == pytest.approx(0.85 * 3.6e6 / stint_s)
    assert rt.ledger_declarations_from_ev_params(None) == {"_provenance": "FSAE_EV_Power_Draw.xlsx"}
    bad = rt.ledger_declarations_from_ev_params({"pack": {"pack_voltage_v": "n/a"},
                                                 "motor": {"pack_voltage_ep_v": 400}})
    assert bad["powertrain"] == {"voltage_v": 400.0}


# --------------------------------------------------- correcting the workbook itself
def _evaluate(b, tmp_path, name="w.xlsx"):
    from suspension.xlsx_formula_cache import evaluate_workbook
    p = tmp_path / name
    p.write_bytes(b)
    vals, report = evaluate_workbook(str(p))
    return vals, report


def _sheet_profile(v_ms, t, n=N):
    """What the sheet holds: the profile resampled to n rows and rounded as written."""
    ts = np.round(rt.resample_profile(t, n), 6)
    vs = np.round(rt.resample_profile(np.asarray(v_ms) * MPH, n), 4) * 0.44704
    a = np.empty(n)
    a[1:-1] = (vs[2:] - vs[:-2]) / (ts[2:] - ts[:-2])
    a[0], a[-1] = (vs[1] - vs[0]) / (ts[1] - ts[0]), (vs[-1] - vs[-2]) / (ts[-1] - ts[-2])
    return ts, vs, a


def test_corrected_workbook_is_a_live_road_load_model(tmp_path):
    v = 10.0 + 5.0 * np.sin(np.linspace(0, 3, N)) ** 2
    t = np.arange(N) * 0.5
    r = rt.lap_to_excel_roundtrip(v, t, _stock(), vehicle={"vehicle_mass_kg": 250.0})
    vals, report = _evaluate(r.excel_bytes, tmp_path)
    assert report["errors"] == {}
    ts, vs, a = _sheet_profile(v, t)
    F = 250.0 * a + 0.5 * RHO * CDA * vs ** 2 + CRR * 250.0 * 9.81
    I = np.maximum(F * vs, 0) / ETA / 504.0
    road = [vals[("KX_RoadLoad", f"F{2 + i}")] for i in range(N)]
    assert road == pytest.approx(I, rel=1e-9)
    assert _load(r.excel_bytes)["KX_RoadLoad"]["K1"].value == 250.0      # the input, editable
    # Every gear column of the current block is that one pack current.
    for col in ("H", "N", "V"):
        assert [vals[("ElecPropulsion", f"{col}{N + 2 + i}")] for i in range(N)] == pytest.approx(I)
    # Motor RPM: wheel rpm times the reduction (gear 7 is column N, 1/7).
    assert vals[("ElecPropulsion", "N2")] == pytest.approx(7 * vs[0] / 0.44704 * K_RPM, rel=1e-9)
    # Pack resistance S R / P, and the joule-heat energy as sum(I^2 R dt).
    assert vals[("Battery Pack Calcs", "B11")] == pytest.approx(140 * 0.012 / 4)
    joule_w = I ** 2 * 0.42
    dt_ = np.diff(ts, prepend=ts[0])
    assert vals[("ThermalLoad", "B6")] == pytest.approx(np.sum(joule_w * dt_) / 3.6e6, rel=1e-9)
    assert vals[("ThermalLoad", "B5")] == pytest.approx(joule_w.max() * 0.42, rel=1e-9)
    assert vals[("ThermalLoad", "B4")] == pytest.approx(joule_w.min() * 0.42, rel=1e-9)
    # EMFs: the pack's terminal voltage under load, V - I R, one row per sample.
    emf = [vals[("EMFs", f"B{1 + i}")] for i in range(N)]
    assert emf == pytest.approx(504.0 - I * 0.42, rel=1e-12)


def test_every_correction_is_listed_with_its_before_and_after():
    r = rt.lap_to_excel_roundtrip(*_const_lap(12.0), _stock())
    wb = _load(r.excel_bytes)
    wc = wb["KX_Corrections"]
    rows = [[c.value for c in row] for row in wc.iter_rows(min_row=2)]
    assert rows == [[f["sheet"], f["cells"], f["before"], f["after"], f["why"]] for f in r.corrections]
    assert all(c.data_type == "s" for row in wc.iter_rows() for c in row if c.value)
    by_cells = {(f["sheet"], f["cells"]): f for f in r.corrections}
    rpm = by_cells[("ElecPropulsion", f"H2:V{N + 1}")]
    assert "*$H$1*1056" in rpm["before"] and "/$H$1*1056" in rpm["after"]
    assert by_cells[("ElecPropulsion", f"H{N + 2}:V{2 * N + 1}")]["after"] == "=KX_RoadLoad!$F2"
    assert by_cells[("ElecPropulsion", f"G{3 * N + 8}")]["after"] == "Min Power (kW)"
    assert by_cells[("ElecPropulsion", f"H{3 * N + 4}:V{3 * N + 4}")]["after"] == f"=MIN(H2:H{N + 1})"
    assert by_cells[("Battery Pack Calcs", "B11")]["after"] == "=B3*B8/B2"
    assert by_cells[("ThermalLoad", "B5")]["after"] == "=MAX(F1:F40)*'Battery Pack Calcs'!B11"
    assert by_cells[("ThermalLoad", "E1")]["after"] == "Joule Heat (W)"
    g = by_cells[("BearingBlowOut", f"G1:G{N}")]["after"]
    assert "TAN(RADIANS($B$17))" in g and "/($B$16/2)" in g
    assert by_cells[("BearingBlowOut", "B19")]["after"] == f"=MAX(G1:G{N})"
    assert ("EMFs", "B1") not in by_cells                       # superseded by the rewrite
    emf = by_cells[("EMFs", f"B1:B{N}")]
    assert emf["before"].endswith(f"ElecPropulsion!H{2 * N + 3}:H{3 * N + 3}*'Battery Pack Calcs'!B11")
    assert emf["after"] == "='Battery Pack Calcs'!$B$12-KX_RoadLoad!$F2*'Battery Pack Calcs'!$B$11"
    assert "V - I R" in emf["why"]
    out = wb["ElecPropulsion"]
    assert [out.cell(row=rr, column=7).value for rr in (2, N + 2, 2 * N + 3)] == \
        ["Motor RPM", "Pack Current (A)", "Pack Power (kW)"]
    assert any("Corrected" in w and "flagged 0 for review" in w for w in r.warnings)


def test_a_second_export_does_not_correct_twice():
    first = rt.lap_to_excel_roundtrip(*_const_lap(12.0), _stock())
    again = rt.lap_to_excel_roundtrip(*_const_lap(12.0), first.excel_bytes)
    assert again.corrections == []
    wb = _load(again.excel_bytes)
    assert wb.sheetnames.count("KX_RoadLoad") == 1 and wb.sheetnames.count("KX_Corrections") == 1
    assert wb["ElecPropulsion"][f"H{N + 2}"].value == "=KX_RoadLoad!$F2"


def test_a_flattened_file_is_rebuilt_corrected_and_live(tmp_path):
    v, t = _const_lap(15.0)
    r = rt.lap_to_excel_roundtrip(v, t, _stock(live=False, svt_name="Speed Vs Time"))
    assert r.ok
    vals, report = _evaluate(r.excel_bytes, tmp_path)
    ep = _load(r.excel_bytes)["ElecPropulsion"]
    assert ep["N2"].value == "=('Speed Vs Time'!B2*1056/$N$1)/($B$10*PI())"
    assert vals[("ElecPropulsion", "N2")] == pytest.approx(7 * 15 * MPH * K_RPM, rel=1e-6)
    I = (0.5 * RHO * CDA * (15 * MPH * 0.44704) ** 2 + CRR * M * 9.81) * 15 * MPH * 0.44704 / ETA / 504
    assert vals[("ElecPropulsion", f"N{N + 2}")] == pytest.approx(I, rel=1e-4)
    assert ep.cell(row=3 * N + 5, column=14).value == f"=MAX(N2:N{N + 1})"
    assert vals[("ElecPropulsion", f"H{3 * N + 6}")] == pytest.approx(               # Min Current
        vals[("KX_RoadLoad", "F2")], rel=1e-9)
    assert any(f["before"].startswith("static values") for f in r.corrections)


def test_emf_columns_need_the_pack_voltage_and_resistance_rows():
    wb = _load(_stock())
    wb["EMFs"]["E1"] = wb["EMFs"]["B1"].value                 # the second EMF column
    buf = io.BytesIO()
    wb.save(buf)
    keep = _load(buf.getvalue())
    keep["EMFs"]["B45"] = "notes below the column"
    b2 = io.BytesIO()
    keep.save(b2)
    out = _load(rt.lap_to_excel_roundtrip(*_const_lap(12.0), b2.getvalue()).excel_bytes)["EMFs"]
    assert out["B45"].value == "notes below the column"     # (pack rows present there: rewritten)
    assert out["E1"].value.startswith("='Battery Pack Calcs'!$B$12-KX_RoadLoad!$F2")
    wb["Battery Pack Calcs"]["A12"] = "Spare"                # no pack nominal voltage row
    buf = io.BytesIO()
    wb.save(buf)
    r = rt.lap_to_excel_roundtrip(*_const_lap(12.0), buf.getvalue())
    cells = {(f["sheet"], f["cells"]): f for f in r.corrections}
    assert cells[("EMFs", "B1")]["why"].startswith("Left for review")
    assert cells[("EMFs", "E1")]["why"].startswith("Left for review")


def test_corrections_leave_alone_what_they_cannot_recognise():
    wb = _load(_stock())
    ep = wb["ElecPropulsion"]
    ep[f"H{N + 2}"] = "=42"                                  # not the current formula
    wb["ThermalLoad"]["B6"] = "=SUM(F1:F7)"                  # does not span the profile
    wb["ThermalLoad"]["A6"] = None
    wb["ElecPropulsion"]["J1"].number_format = "h:mm:ss"     # a date-formatted ratio
    ep[f"I{3 * N + 4}"] = 5                                  # a typed number in a summary row
    wb["Battery Pack Calcs"]["A11"] = "Spare"                # no pack-resistance row to fix
    wb["BearingBlowOut"]["G3"] = None
    buf = io.BytesIO()
    wb.save(buf)
    r = rt.lap_to_excel_roundtrip(*_const_lap(12.0), buf.getvalue())
    assert any(f"ElecPropulsion!H{N + 2} does not hold the current formula" in w for w in r.warnings)
    cells = {(f["sheet"], f["cells"]): f for f in r.corrections}
    assert cells[("ThermalLoad", "B6")]["why"].startswith("Left for review")
    assert cells[("ElecPropulsion", "J1")]["after"] == "General"
    assert ("ElecPropulsion", f"H{N + 2}:V{2 * N + 1}") not in cells
    assert not any(f["sheet"] == "Battery Pack Calcs" for f in r.corrections)
    assert cells[("BearingBlowOut", f"G1:G{N}")]["after"].count("RADIANS") == 1
    out = _load(r.excel_bytes)["ElecPropulsion"]
    assert out[f"H{N + 2}"].value == "=42" and out[f"I{N + 2}"].value == f"=($B$8*$B$11*I2)/1000"


def test_rebuilt_blocks_place_by_contract_or_say_why_not():
    wb = _load(_stock(live=False, dependents=False))
    ws = wb.create_sheet("Readers")
    ws["A1"] = f"=SUM(ElecPropulsion!H2:V{N + 1})"
    ws["A2"] = f"=SUM(ElecPropulsion!H10:V{N + 9})"               # overlaps the first block
    buf = io.BytesIO()
    wb.save(buf)
    r = rt.lap_to_excel_roundtrip(*_const_lap(12.0), buf.getvalue())
    assert any("do not fit 40-row blocks" in w and "stock spacing (2, 42, 83)" in w for w in r.warnings)


def test_helpers_for_the_corrections():
    blocks = [(2, 41), (42, 81), (83, 122)]
    assert rt._fix_block_ranges("=MIN(H2:H42)+ElecPropulsion!H2:H42", blocks, own_sheet=True) == \
        "=MIN(H2:H41)+ElecPropulsion!H2:H42"
    assert rt._fix_block_ranges("=SUM(ElecPropulsion!$H$83:$V$123)", blocks, own_sheet=False) == \
        "=SUM(ElecPropulsion!$H$83:$V$122)"
    assert rt._fix_block_ranges("=SUM(ElecPropulsion!H5:H9)", blocks, own_sheet=False) == \
        "=SUM(ElecPropulsion!H5:H9)"
    assert rt._wrap_tan_in_radians("=TAN(A1)+tan(B2*(1+2))") == \
        "=TAN(RADIANS(A1))+TAN(RADIANS(B2*(1+2)))"
    assert rt._wrap_tan_in_radians("=TAN(") == "=TAN("
    assert rt._formula_text(ArrayFormula("A1:A2", "=X")) == "=X" and rt._formula_text(3) == 3


@pytest.fixture(scope="module")
def stock_layout():
    """The stock file's addresses: 1893 profile rows, BatteryPackConfig, cached gear ratios."""
    return _stock(n=1893, numeric_ratios=True, pack_name="BatteryPackConfig")


def test_the_audit_finds_nothing_left_to_fix_in_a_corrected_file(stock_layout, tmp_path):
    from suspension import power_draw as pd
    before_p, after_p = tmp_path / "before.xlsx", tmp_path / "after.xlsx"
    before_p.write_bytes(stock_layout)
    found = {f.check for f in pd.audit_workbook(str(before_p))}
    assert found >= {"pd-no-vehicle-model", "pd-gear-inverted", "pd-efficiency-and-phase",
                     "pd-rpm-stat-range", "pd-power-stat-units", "pd-minmax-inconsistent",
                     "pd-sum-of-power", "pd-thermal-units", "pd-emf-dimensions", "pd-emf-range",
                     "pd-tan-degrees", "pd-pitch-diameter", "pd-bearing-max-range"}
    t = np.arange(1893) / 15.0
    v = 12.0 + 4.0 * np.sin(t / 7.0) ** 2
    r = rt.lap_to_excel_roundtrip(v, t, stock_layout)
    assert r.ok and not any("Resampled" in w for w in r.warnings)
    after_p.write_bytes(r.excel_bytes)
    left = {f.check for f in pd.audit_workbook(str(after_p))}
    assert left == set()
    readers = {f.check for f in pd.audit_kinematik_readers(str(after_p))}
    assert "pd-roundtrip-mirrors-errors" not in readers and "pd-reader-sheet-name" not in readers


def test_an_unlabelled_energy_cell_is_converted_but_not_relabelled():
    wb = _load(_stock())
    wb["ThermalLoad"]["A6"] = None
    buf = io.BytesIO()
    wb.save(buf)
    out = _load(rt.lap_to_excel_roundtrip(*_const_lap(12.0), buf.getvalue()).excel_bytes)
    assert out["ThermalLoad"]["B6"].value.startswith("=SUMPRODUCT(") and out["ThermalLoad"]["A6"].value is None


def test_the_reader_audit_checks_behaviour(tmp_path, monkeypatch):
    from suspension import power_draw as pd
    p = tmp_path / "nopack.xlsx"
    p.write_bytes(_stock(pack=False))
    assert "pd-reader-sheet-name" in {f.check for f in pd.audit_kinematik_readers(str(p))}

    flat = rt.ExcelRoundTripResult(current_draw_a=np.ones(40))
    monkeypatch.setattr(rt, "lap_to_excel_roundtrip_from_db", lambda *a, **k: flat)
    assert "pd-roundtrip-mirrors-errors" in {f.check for f in pd.audit_kinematik_readers(str(p))}
