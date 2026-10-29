"""Characterization suite for suspension/ev_electrical_check.py (mutation gate); see tests/_golden.py."""
import sys

import numpy as np
import openpyxl
import pytest

from suspension import ev_electrical_check as ec
from tests._golden import check, plain

PACK = [("Parrallel Battery Count", 4), ("Series Battery Count", 120), ("Capacity Battery Cell (Ah)", 4.5),
        ("Internal Resistance Battery Cell (Ohms)", 0.02), ("Battery Pack Nominal Voltage (V)", 430.0),
        ("Pack Capacity (Ah)", 18.0), ("Fuse Max (A)", 60.0), ("Current Across a Single Cell (A)", 14.0)]
EP = [("Motor Peak Power (kW)", 80.0), ("Motor Peak Torque (Nm)", 200.0), ("Motor Efficiency ", 0.9),
      ("Motor Max DC Voltage (V)", 600.0), ("Wheel Diameter (in)", 16.0)]


def _wb(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    if not wb.worksheets:
        wb.create_sheet("x")
    wb.save(path)
    return str(path)


def _ep_rows(extra_at=None):
    rows = [("pad", None)] * 6 + EP
    if extra_at:
        rows = rows + [("pad", None)] * (extra_at - len(rows) - 1) + [("Wheel Diameter (in)", 99.0)]
    return rows


def test_from_excel_cases(tmp_path, monkeypatch):
    cases = {
        "full": {"BatteryPackConfig": PACK + [(None, 1), ("blank", None), ("  Fuse Max (A)  ", 70.0)],
                 "ElecPropulsion": EP},
        "aliases": {"BatteryPackConfig": [("parallel battery count", 2), (" series battery count ", 100),
                                           ("Fuse Max (A)", "abc")],
                    "ElecPropulsion": [("motor efficiency", 0.8)]},
        "defaults": {"Battery Pack Calcs": [("Fuse Max (A)", 0)], "Other": [("a", 1)]},
        "one_col": {"BatteryPackConfig": [("Series Battery Count",), ("Fuse Max (A)",)],
                    "ElecPropulsion": [("Motor Peak Power (kW)",)]},
        "ep_row12": {"BatteryPackConfig": PACK, "ElecPropulsion": _ep_rows(12)},
        "ep_row13": {"BatteryPackConfig": PACK, "ElecPropulsion": _ep_rows(13)},
        "ep_late": {"BatteryPackConfig": PACK, "ElecPropulsion": [("pad", 0)] * 11 + EP},
        "zero_par": {"BatteryPackConfig": [("Parrallel Battery Count", 0.4), ("Fuse Max (A)", 30.0)]},
        "empty": {},
    }
    res = {}
    for k, sheets in cases.items():
        p = _wb(tmp_path / f"{k}.xlsx", sheets)
        res[k] = ec.ElecParams.from_excel(p)
        res[k + "_b"] = ec.ElecParams.from_excel_bytes(open(p, "rb").read())
    res["missing"] = ec.ElecParams.from_excel(str(tmp_path / "nope.xlsx"))
    res["garbage"] = ec.ElecParams.from_excel_bytes(b"not a workbook")
    res["direct"] = ec.ElecParams(pack_voltage_v=300.0, pack_capacity_ah=8.0, fuse_max_a=40.0)
    check("ev_elec_params", res)
    p = _wb(tmp_path / "full2.xlsx", cases["full"])
    data = open(p, "rb").read()
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    assert plain(ec.ElecParams.from_excel(p)) == plain(ec.ElecParams())
    assert plain(ec.ElecParams.from_excel_bytes(data)) == plain(ec.ElecParams())
    with pytest.raises(ImportError, match="^openpyxl is required to read Excel files. Install it with: pip"):
        ec.load_speed_vs_time(p)


def test_power_helpers():
    res = {}
    v = np.array([0.0, 5.0, 20.0, 35.0])
    a = np.array([3.0, -8.0, 0.5, 0.0])
    res["tp"] = [ec.traction_power_w(v, a), ec.traction_power_w(v, a, vehicle_mass_kg=300.0, drag_cda=0.9,
                                                                crr=0.02, rho_air=1.1, drivetrain_eff=0.3),
                 ec.traction_power_w(12.0, 1.0, drivetrain_eff=0.8)]
    res["fls"] = [ec.fuse_limited_speed(p) for p in (0.0, 5e3, 25e3, 1e8)]
    res["fls_kw"] = [ec.fuse_limited_speed(25e3, vehicle_mass_kg=300.0, drag_cda=0.8, crr=0.03, rho_air=1.0,
                                           drivetrain_eff=0.95),
                     ec.fuse_limited_speed(25e3, drivetrain_eff=0.2)]
    check("ev_elec_power", res)


def _lap(n=60):
    s = np.linspace(0.0, 600.0, n)
    v = 12.0 + 10.0 * np.sin(s / 60.0) ** 2
    return v, s


def test_lap_checks():
    v, s = _lap()
    P = ec.ElecParams
    cases = {
        "default": (v, s, P(), {}),
        "fuse": (v * 2.2, s, P(), {}),
        "cell": (v, s, P(nominal_cell_current_a=5.0), {}),
        "energy": (v, s * 40.0, P(pack_capacity_ah=1.0), {}),
        "all": (v * 2.5, s * 30.0, P(pack_capacity_ah=0.5, nominal_cell_current_a=2.0), {}),
        "small_pack": (v, s, P(pack_capacity_ah=9.9), {}),
        "cap10": (v, s, P(pack_capacity_ah=10.0), {}),
        "stopped": (np.zeros(5), np.linspace(0, 4, 5), P(pack_capacity_ah=0.0), {}),
        "crawl": (np.array([0.0, 0.2, 0.4, 0.3, 0.0]), np.array([0.0, 0.1, 0.3, 0.5, 0.6]), P(), {}),
        "lowv": (v, s, P(pack_voltage_v=0.5, n_parallel=0), {}),
        "kw": (v, s, P(fuse_max_a=20.0), dict(drivetrain_eff=0.3, vehicle_mass_kg=320.0, drag_cda=1.4,
                                              crr=0.03, rho_air=1.0)),
        "short": ([5.0], [0.0], P(), {}),
        "empty": ([], [], P(), {}),
        "coast": (np.linspace(30.0, 10.0, 20), np.linspace(0.0, 300.0, 20), P(), {}),
    }
    res = {k: ec.check_lap_electrical(a, b, p, **kw) for k, (a, b, p, kw) in cases.items()}
    t = np.linspace(0.0, 30.0, 31)
    mph = 20.0 + 15.0 * np.sin(t / 5.0)
    res["csv"] = ec.check_lap_from_speed_csv(mph, t, P())
    res["csv_kw"] = ec.check_lap_from_speed_csv(mph, t, P(), drivetrain_eff=0.7, vehicle_mass_kg=280.0,
                                                drag_cda=1.3, crr=0.025)
    res["csv_one"] = ec.check_lap_from_speed_csv([3.0], [0.0], P())
    check("ev_elec_lap", res)


def test_speed_vs_time_loader(tmp_path):
    res = {}
    rows = [("Time (s)", "Speed (mph)"), (0, 10), (1, 12.5), ("x", 3), (2, None), (3.5, 20)]
    ok = _wb(tmp_path / "ok.xlsx", {"SpeedVsTime": rows})
    res["ok"] = ec.load_speed_vs_time(ok)
    res["ok_b"] = ec.load_speed_vs_time_from_bytes(open(ok, "rb").read())
    res["alias"] = ec.load_speed_vs_time(_wb(tmp_path / "al.xlsx", {"Speed Vs Time": [(0, 1), (1, 2)]}))
    one = _wb(tmp_path / "one.xlsx", {"SpeedVsTime": [(0,), (1,), (2,)]})
    few = _wb(tmp_path / "few.xlsx", {"SpeedVsTime": [(0, 1), ("a", "b")]})
    none = _wb(tmp_path / "none.xlsx", {"A": [(1, 2)], "B": [(1, 2)]})
    for k, p in (("one", one), ("few", few), ("none", none), ("bad", str(tmp_path / "zz.xlsx"))):
        try:
            ec.load_speed_vs_time(p)
            res[k] = "no error"
        except ValueError as e:
            res[k] = str(e).replace(str(tmp_path), "<T>")
    try:
        ec.load_speed_vs_time_from_bytes(open(none, "rb").read())
    except ValueError as e:
        res["none_b"] = str(e)
    check("ev_elec_svt", res)


def test_more_lap_edges():
    P = ec.ElecParams
    res = {"two": ec.check_lap_electrical([10.0, 12.0], [0.0, 20.0], P()),
           "tiny_pack": ec.check_lap_electrical(np.linspace(5, 8, 6), np.linspace(0, 50, 6), P(pack_capacity_ah=1.0)),
           "csv_half": ec.check_lap_from_speed_csv([10.0, 14.0, 20.0, 18.0, 12.0], [0.0, 0.5, 1.0, 1.5, 2.0], P())}
    check("ev_elec_edges", res)
