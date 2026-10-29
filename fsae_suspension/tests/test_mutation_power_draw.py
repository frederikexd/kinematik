"""Characterization suite for suspension/power_draw.py (mutation gate).

Each case is a purpose-built workbook or trace; the full findings (check, severity, message, cells, detail) and the
numeric traces are compared against tests/golden/power_draw_*.json. See tests/_golden.py."""
import datetime as dt
import math
import json
import re

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from suspension import power_draw as pdw
from suspension.interfaces import IntegrationLedger, SubsystemInterface
from tests._golden import check, plain

# every value differs from the PackSpec default, so a reader that drops a field is seen
PACK_ROWS = [("Fuse Max (A)", 60), ("Parrallel Battery Count", 4), ("Series Battery Count", 120),
             ("Nominal Battery Voltage (V)", 3.7), ("Capacity Battery Cell (Ah)", 4.5),
             ("Endurance Length (km)", 22), ("Max Battery Cells", 600),
             ("Internal Resistance Battery Cell (Ohms)", 0.015), ("Battery Cell Weight (kg)", 0.07)]


def _cache(path, cached):
    """Give formula cells the cached result Excel would have stored ({formula_text: value})."""
    import shutil
    import zipfile
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                txt = data.decode()
                for f, v in cached.items():
                    txt = txt.replace(f"<f>{f}</f><v></v>", f"<f>{f}</f><v>{v}</v>")
                data = txt.encode()
            zout.writestr(item, data)
    shutil.move(tmp, path)
    return path


def _norm(x, tmp_path):
    s = json.dumps(plain(x)).replace(str(tmp_path), "<T>")
    return json.loads(re.sub(r"/tmp/tmp\w+\.xlsx", "<TMP>", s))


def _save(tmp_path, name, sheets, pack=None, first=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if first:
        wb.create_sheet(first)
    if pack is not None:
        ws = wb.create_sheet("BatteryPackConfig")
        for i, (a, b) in enumerate(pack, start=1):
            ws.cell(i, 1, a)
            ws.cell(i, 2, b)
    for sname, cells in sheets.items():
        ws = wb.create_sheet(sname)
        for coord, val in cells.items():
            ws[coord] = val
    if not wb.worksheets:
        wb.create_sheet("Empty")
    p = tmp_path / f"{name}.xlsx"
    wb.save(p)
    return str(p)


BASE = {
    "ElecPropulsion": {"H1895": "=B$11*H2", "H2": "=$A2*$H$1*1056", "N1": 0.125, "H3789": "=SQRT(3)*H1895",
                       "H1": dt.datetime(2026, 1, 8), "V1": dt.datetime(2026, 1, 10), "J1": 3,
                       "H5683": "=MIN(H2:H1895)", "G5687": "Min (W)", "G5688": "Max (kW)"},
    "ThermalLoad": {"B4": "=MIN(C1:C9)*B2", "B5": "=MAX(C1:C9)*B1*B2", "B6": "=SUM(C1:C9)",
                    "E1": "Heat (kW)", "F1": "=B11*C1"},
    "EMFs": {"B1": "=B12*B11*C5682"},
    "BearingBlowOut": {"G1": "=A1/$B$16*TAN(20)", "B19": "=MAX(G1:G1894)"},
    "Sheet1": {"A19": 1.5, "A21": 2, "A23": 2.5, "A25": 3, "A27": 3.5, "A29": 4, "A31": 3.0, "A20": 9},
    "Data Dashboard": {},
}


def _variant(**edits):
    out = {k: dict(v) for k, v in BASE.items()}
    for key, val in edits.items():
        sheet, coord = key.split("__")
        sheet = sheet.replace("_", " ") if sheet == "Data_Dashboard" else sheet
        if val is None:
            out[sheet].pop(coord, None)
        else:
            out[sheet][coord] = val
    return out


WB_CASES = {
    "base": _variant(),
    "h1895_no_b11": _variant(ElecPropulsion__H1895="=B11*H2"),
    "h1895_no_h2": _variant(ElecPropulsion__H1895="=B$11*H3"),
    "h1895_sqrt": _variant(ElecPropulsion__H1895="=sqrt(B$11*H2)", ElecPropulsion__H3789="=1"),
    "h1895_array": _variant(ElecPropulsion__H1895=ArrayFormula("H1895", "=B$11*H2")),
    "h2_digit": _variant(ElecPropulsion__H2="=$A2*$H$10*1056"),
    "h2_space": _variant(ElecPropulsion__H2="=$A2* $H$1*1056"),
    "h2_no1056": _variant(ElecPropulsion__H2="=$A2*$H$1*1000"),
    "n1_one": _variant(ElecPropulsion__N1=1),
    "n1_zero": _variant(ElecPropulsion__N1=0),
    "n1_text": _variant(ElecPropulsion__N1="0.5"),
    "n1_big": _variant(ElecPropulsion__N1=7),
    "h3789_lower": _variant(ElecPropulsion__H3789="=sqrt(3)*x"),
    "dates_edges": _variant(ElecPropulsion__H1=None, ElecPropulsion__V1=None,
                            ElecPropulsion__G1=dt.datetime(2026, 1, 1), ElecPropulsion__W1=dt.datetime(2026, 1, 2),
                            ElecPropulsion__I1=dt.time(3, 0)),
    "h5683_other": _variant(ElecPropulsion__H5683="=MIN(H2:H1894)"),
    "stat_lo_only": _variant(ElecPropulsion__G5688="Max (W)"),
    "stat_hi_only": _variant(ElecPropulsion__G5687="Min (kW)"),
    "b4_empty": _variant(ThermalLoad__B4=None),
    "b5_empty": _variant(ThermalLoad__B5=None),
    "b5_b10": _variant(ThermalLoad__B5="=MAX(C1:C9)*B10*B2"),
    "b4_b1": _variant(ThermalLoad__B4="=MIN(C1:C9)*B1*B2"),
    "b6_sumproduct": _variant(ThermalLoad__B6="=SUMPRODUCT(C1:C9)"),
    "b6_lower": _variant(ThermalLoad__B6="=  sum(C1:C9)"),
    "e1_w": _variant(ThermalLoad__E1="Heat (W)"),
    "f1_b12": _variant(ThermalLoad__F1="=B12*C1"),
    "emf_no_b12": _variant(EMFs__B1="=B11*C5682"),
    "emf_no_b11": _variant(EMFs__B1="=B12*C1"),
    "emf_kx": _variant(EMFs__B1="=B12*B11*KX_RoadLoad"),
    "tan_lower": _variant(BearingBlowOut__G1="=A1/$B$16/2*tan(20)"),
    "tan_radians": _variant(BearingBlowOut__G1="=A1*tan(radians(20))"),
    "b19_other": _variant(BearingBlowOut__B19="=MAX(G1:G1893)"),
    "sheet1_text": _variant(Sheet1__A19="x", Sheet1__A21=None, Sheet1__A31=None),
    "sheet1_some": _variant(Sheet1__A23=None, Sheet1__A27=None, Sheet1__A29=None),
    "n1_cached": _variant(ElecPropulsion__N1="=1/8"),
    "dash_full": _variant(**{"Data_Dashboard__A2": 1}),
}


def test_audit_workbook_cases(tmp_path):
    res = {}
    for name, sheets in WB_CASES.items():
        res[name] = pdw.audit_workbook(_cache(_save(tmp_path, name, sheets), {"1/8": 0.125}))
    res["no_sheets"] = pdw.audit_workbook(_save(tmp_path, "none", {}, pack=PACK_ROWS))
    check("power_draw_workbook", _norm(res, tmp_path))


def test_numeric_label_and_sheet_helpers(tmp_path):
    res = {}
    for i, v in enumerate([None, True, dt.datetime(2020, 1, 1), dt.date(2020, 1, 1), dt.time(1, 2),
                           dt.timedelta(1), " 3.5 ", "abc", 4, 2.5]):
        try:
            res[f"num{i}"] = pdw._numeric(v, "S!A1")
        except pdw.WorkbookReadError as e:
            res[f"num{i}"] = str(e)
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, (a, b, c) in enumerate([("Alpha", 1, 10), ("  ", 2, 20), (None, 3, 30), (5, 4, 40),
                                   (" alpha ", 5, 50), ("Beta", 6, 60)], start=1):
        ws.cell(r, 1, a); ws.cell(r, 2, b); ws.cell(r, 3, c)
    res["lm"] = pdw._label_map(ws)
    res["lm_c3"] = pdw._label_map(ws, value_col=3)
    res["lm_max5"] = pdw._label_map(ws, max_row=5)
    res["lm_lab2"] = pdw._label_map(ws, label_col=2, value_col=3)
    w2 = openpyxl.Workbook()
    a = w2.active
    a.cell(40, 1, "Fuse Max")
    b = w2.create_sheet("b")
    b.cell(41, 1, "Series Battery Count")
    b.cell(1, 1, "fuse max")
    res["fs1"] = getattr(pdw._find_sheet(w2, "FUSE max"), "title", None)
    res["fs2"] = getattr(pdw._find_sheet(w2, "series battery count"), "title", None)
    res["fs3"] = getattr(pdw._find_sheet(w2, "fuse max", "zzz"), "title", None)
    res["fs4"] = getattr(pdw._find_sheet(w2), "title", None)
    check("power_draw_helpers", _norm(res, tmp_path))


def _pack_variants():
    alias = [(a.replace("Parrallel", "Parallel"), b) for a, b in PACK_ROWS]
    no_par = [r for r in PACK_ROWS if "allel" not in r[0]]
    series_only = [r for r in PACK_ROWS if "Fuse" not in r[0]]
    fuse_only = [r for r in PACK_ROWS if "Series" not in r[0]]
    bad_cell = [(a, dt.datetime(2020, 1, 1) if "Nominal" in a else b) for a, b in PACK_ROWS]
    text_cell = [(a, "3.7" if "Nominal" in a else b) for a, b in PACK_ROWS]
    other = [("Series Battery Count", 999)] + [("x", 1)]
    series_alone = [r for r in PACK_ROWS if "Fuse" not in r[0] and "allel" not in r[0]]
    cached = [(a, "=60+0" if "Fuse" in a else b) for a, b in PACK_ROWS]
    return {"plain": PACK_ROWS, "alias": alias, "no_par": no_par, "series_only": series_only,
            "fuse_only": fuse_only, "bad_cell": bad_cell, "text_cell": text_cell, "other_first": other,
            "unlabelled": [("nothing", 1)], "series_alone": series_alone, "cached": cached,
            "decoy_par": PACK_ROWS}


def test_read_pack_config_cases(tmp_path):
    res = {}
    for name, rows in _pack_variants().items():
        first = None
        if name in ("other_first", "decoy_par"):
            p = _save(tmp_path, name, {}, pack=PACK_ROWS, first="Decoy")
            wb = openpyxl.load_workbook(p)
            wb["Decoy"].cell(1, 1, "Series Battery Count" if name == "other_first" else "Parrallel Battery Count")
            wb["Decoy"].cell(1, 2, 999)
            wb.save(p)
        else:
            p = _cache(_save(tmp_path, name, {}, pack=rows, first=first), {"60+0": 60})
        try:
            res[name] = pdw.read_pack_config(p)
        except pdw.WorkbookReadError as e:
            res[name] = str(e)
    check("power_draw_read", _norm(res, tmp_path))


def test_reader_probe_cases(tmp_path):
    p = _save(tmp_path, "probe", {}, pack=PACK_ROWS, first="Front")

    def reads(path):
        ws = openpyxl.load_workbook(path)["BatteryPackConfig"]
        return [ws["B2"].value, ws["B3"].value, ws["B4"].value]

    def front(path):
        ws = openpyxl.load_workbook(path)["Front"]
        return [ws["B2"].value, ws["B3"].value]

    def boom(path):
        raise KeyError("nope")

    class Fixed:
        def __call__(self, path):
            return 1

    res = {"reads": pdw.assert_reader_sees_the_file(p, reads),
           "ignores": pdw.assert_reader_sees_the_file(p, lambda path: 7),
           "boom": pdw.assert_reader_sees_the_file(p, boom),
           "probe_b4": pdw.assert_reader_sees_the_file(p, reads, probe={"B4": 1.0}),
           "probe_b9": pdw.assert_reader_sees_the_file(p, reads, probe={"B9": 1.0}),
           "fixed_obj": pdw.assert_reader_sees_the_file(p, Fixed()),
           "pack": pdw.assert_reader_sees_the_file(p, pdw.read_pack_config)}
    q = _save(tmp_path, "probe2", {"Front": {"B2": 1}}, pack=None)
    res["front"] = pdw.assert_reader_sees_the_file(q, front)
    out = _norm(res, tmp_path)
    for k in ("ignores", "fixed_obj"):
        for f in out[k]:
            f["message"] = re.sub(r"<function .*?>|<.*?Fixed object at 0x\w+>", "<obj>", f["message"])
    check("power_draw_probe", out)


PACKS = [pdw.PackSpec(), pdw.PackSpec(n_series=100, n_parallel=2, cell_voltage_v=4.0, cell_capacity_ah=3.0,
                                      cell_resistance_ohm=0.02, cell_weight_kg=0.05, fuse_max_a=80.0,
                                      max_cells=150),
         pdw.PackSpec(n_series=10, n_parallel=1, cell_resistance_ohm=0.5, max_cells=None)]


def test_pack_spec_and_gear_numbers():
    res = []
    for p in PACKS:
        row = {m: getattr(p, m)() for m in ("cell_count", "nominal_voltage_v", "capacity_ah", "energy_kwh",
                                            "mass_kg", "resistance_ohm", "max_deliverable_power_w",
                                            "power_at_fuse_limit_w")}
        row["cfp"] = [p.current_for_power(w) for w in (-5.0, 0.0, 1.0, 2e4, p.max_deliverable_power_w(),
                                                       p.max_deliverable_power_w() * 1.01)]
        row["heat"] = p.joule_heat_w(30.0)
        row["cell"] = p.cell_current(30.0)
        res.append(row)
    v = pdw.VehicleSpec(wheel_diameter_in=16.0, drivetrain_efficiency=0.9, motor_efficiency=0.8,
                        inverter_efficiency=0.7)
    d = pdw.DriveSpec(reduction=4.0)
    res.append({"r": v.wheel_radius_m(), "eta": v.total_efficiency(), "rpm": d.motor_rpm(30.0, 16.0),
                "tq": d.motor_torque_nm(500.0, 0.2), "tq2": d.motor_torque_nm(500.0, 0.2, 0.8),
                "ft": pdw.gear_tangential_force_n(50.0, 0.04), "fr": pdw.gear_radial_force_n(100.0),
                "fr30": pdw.gear_radial_force_n(100.0, 30.0)})
    for bad in (0.0, -1.0):
        with pytest.raises(ValueError, match="^pitch diameter must be positive$"):
            pdw.gear_tangential_force_n(1.0, bad)
    check("power_draw_numbers", {str(i): r for i, r in enumerate(res)})


SPEEDS = {
    "launch": [0.0, 2.0, 5.0, 9.0, 14.0, 20.0, 25.0, 28.0, 30.0, 30.0],
    "brake": [40.0, 38.0, 35.0, 30.0, 26.0, 22.0, 20.0, 20.0],
    "glitch": [20.0, 20.0, 60.0, 20.0, 20.0, 21.0, 22.0, 22.0],
    "crawl": [0.0, 0.05, 0.1, 0.2, 0.11, 0.0],
    "fast": [80.0, 85.0, 90.0, 95.0, 100.0, 104.0],
}
VEHICLES = {"default": pdw.VehicleSpec(), "regen": pdw.VehicleSpec(regen=True, regen_efficiency=0.5),
            "light": pdw.VehicleSpec(mass_kg=200.0, crr=0.015, cda_m2=0.8, air_density=1.1,
                                     wheel_diameter_in=16.0, mu_lon=1.1)}
DRIVES = {"d7": pdw.DriveSpec(), "d2": pdw.DriveSpec(reduction=2.0, motor_continuous_torque_nm=10.0),
          "d12": pdw.DriveSpec(reduction=12.0, motor_max_rpm=9000.0)}
TINY = pdw.PackSpec(n_series=20, n_parallel=1, cell_resistance_ohm=0.2, fuse_max_a=20.0, cell_capacity_ah=0.5)


def _trace_cases():
    yield "launch", SPEEDS["launch"], 0.5, pdw.PackSpec(), "default", "d7", 1, True
    yield "launch_regen", SPEEDS["launch"], 0.5, pdw.PackSpec(), "regen", "d2", 1, True
    yield "brake", SPEEDS["brake"], 0.5, pdw.PackSpec(), "default", "d7", 1, True
    yield "brake_regen", SPEEDS["brake"], 0.5, pdw.PackSpec(), "regen", "d12", 1, True
    yield "glitch", SPEEDS["glitch"], 0.1, pdw.PackSpec(), "default", "d7", 1, True
    yield "glitch_s3", SPEEDS["glitch"], 0.1, pdw.PackSpec(), "light", "d7", 3, True
    yield "glitch_s4", SPEEDS["glitch"], 0.1, pdw.PackSpec(), "light", "d2", 4, True
    yield "glitch_noclamp", SPEEDS["glitch"], 0.1, pdw.PackSpec(), "default", "d2", 1, False
    yield "crawl", SPEEDS["crawl"], 1.0, pdw.PackSpec(), "light", "d7", 1, True
    yield "fast_tiny", SPEEDS["fast"], 1.0, TINY, "default", "d12", 1, True
    yield "fast", SPEEDS["fast"], 1.0, PACKS[1], "light", "d12", 2, True
    yield "launch_tiny", [0.0, 30.0], 2.0, TINY, "default", "d2", 1, True


def test_power_draw_trace_and_findings():
    res = {}
    for name, sp, dts, pack, vn, dn, sw, clamp in _trace_cases():
        tr = pdw.power_draw_trace(sp, dts, pack, VEHICLES[vn], DRIVES[dn], smooth_window=sw, clamp_to_grip=clamp)
        res[name] = {"trace": tr, "findings": pdw.trace_findings(tr),
                     "m": [tr.peak_current_a(), tr.peak_power_kw(), tr.energy_kwh(), tr.joule_energy_kwh(),
                           tr.mean_joule_w(), tr.peak_motor_rpm(), tr.peak_motor_torque_nm(), tr.duration_s()]}
    res["disc"] = [pdw.find_discontinuities(SPEEDS["glitch"], 0.1, 300.0),
                   pdw.find_discontinuities(SPEEDS["glitch"], 0.1, 300.0, mu_lon=60.0),
                   pdw.find_discontinuities([0.0, 1.0, 0.0], 0.03, 300.0),
                   pdw.find_discontinuities([0.0, 1.4 * pdw.G * 0.1 / pdw.MPH_TO_MS, 0.0], 0.1),
                   pdw.find_discontinuities([0.0, 1.45 * pdw.G * 0.1 / pdw.MPH_TO_MS], 0.1, None, 1.4)]
    res["defaults"] = pdw.power_draw_trace(SPEEDS["glitch"], 0.1, pdw.PackSpec(), VEHICLES["light"], DRIVES["d7"])
    odd = pdw.DriveSpec(reduction=5.0, motor_max_rpm=15000.0, motor_peak_torque_nm=80.0, motor_peak_power_kw=60.0,
                        motor_continuous_torque_nm=40.0)
    res["sweep_odd"] = {k: [t.drive, t.peak_current_a(), t.peak_motor_torque_nm(), t.clamped_samples, t.smooth_window]
                        for k, t in pdw.gear_sweep(SPEEDS["glitch"], 0.1, pdw.PackSpec(), VEHICLES["light"], odd,
                                                   reductions=(2.0, 6.0)).items()}
    res["ma"] = [pdw._moving_average([1.0, 2.0, 6.0, 3.0, 9.0], w) for w in (0, 1, 2, 3, 4, 5, 9)]
    res["sweep"] = {k: [t.peak_motor_rpm(), t.peak_motor_torque_nm(), t.drive.reduction,
                        t.drive.motor_max_rpm, t.drive.motor_peak_torque_nm, t.drive.motor_peak_power_kw,
                        t.drive.motor_continuous_torque_nm, t.smooth_window, t.clamped_samples]
                    for k, t in pdw.gear_sweep(SPEEDS["glitch"], 0.1, pdw.PackSpec(), VEHICLES["light"],
                                               DRIVES["d12"], reductions=(3, 9.5), smooth_window=3,
                                               clamp_to_grip=False).items()}
    full = pdw.gear_sweep(SPEEDS["launch"], 0.5, pdw.PackSpec(), VEHICLES["default"], DRIVES["d2"])
    res["sweep_keys"] = [[k, full[k].drive.reduction, full[k].clamped_samples] for k in full]
    empty = pdw.PowerDrawTrace(0.1, [], [], [], [], [], [], [], [], [], [], [], 0, 0, [], 1,
                               pdw.PackSpec(), pdw.VehicleSpec(), pdw.DriveSpec())
    res["empty"] = [empty.peak_current_a(), empty.peak_power_kw(), empty.mean_joule_w(),
                    empty.peak_motor_rpm(), empty.peak_motor_torque_nm(), empty.duration_s()]
    for args in ((0.0,), (-1.0,)):
        with pytest.raises(ValueError, match="^dt_s must be positive$"):
            pdw.power_draw_trace([1.0, 2.0], args[0], pdw.PackSpec(), pdw.VehicleSpec(), pdw.DriveSpec())
    with pytest.raises(ValueError, match="^a trace needs at least two samples to differentiate$"):
        pdw.power_draw_trace([1.0], 0.1, pdw.PackSpec(), pdw.VehicleSpec(), pdw.DriveSpec())
    check("power_draw_trace", res)


def test_pack_level_findings():
    res = {}
    for i, p in enumerate(PACKS + [TINY, pdw.PackSpec(n_parallel=4, max_cells=560)]):
        res[f"bms{i}"] = pdw.bms_findings(p)
        res[f"bms16_{i}"] = pdw.bms_findings(p, cells_per_board=12)
        res[f"stack{i}"] = repr(pdw.to_bq_stack(p))
        res[f"stack_p{i}"] = repr(pdw.to_bq_stack(p, part="BQ79656-Q1", cells_per_board=14))
        for j, d in enumerate([pdw.DriveSpec(), pdw.DriveSpec(motor_peak_power_kw=10.0),
                               pdw.DriveSpec(motor_peak_power_kw=1e4)]):
            res[f"mf{i}_{j}"] = pdw.motor_feasibility(p, d)
        res[f"elec{i}"] = repr(pdw.to_elec_params(p, pdw.VehicleSpec()))
        res[f"iface{i}"] = pdw.to_interface(p)
    res["iface_kw"] = pdw.to_interface(pdw.PackSpec(source="src.xlsx!S"), name="hv", heat_reject_w=12.0,
                                       is_estimate=True)
    res["part_name"] = [pdw.part_name(type("D", (), {"part": "X"})()), pdw.part_name(3)]
    sv = [{}, {"pack_resistance_ohm": 1.792}, {"pack_resistance_ohm": pdw.PackSpec().resistance_ohm() + 5e-7},
          {"pack_resistance_ohm": pdw.PackSpec().resistance_ohm() - 2e-6}, {"joule_max_label": "Joule (KWh)"},
          {"joule_max_label": "joule (w)"}]
    res["aps"] = [pdw.audit_pack_sheet(pdw.PackSpec(), s) for s in sv]
    check("power_draw_pack", res)


def test_sync_ledger_cases():
    res = {}
    pack = pdw.PackSpec(source="wb.xlsx!P")
    fresh = pdw.to_interface(pack)
    led = IntegrationLedger()
    res["empty"] = [pdw.sync_ledger(led, pack), led.interfaces["electrics"]]
    led = IntegrationLedger()
    led.set(SubsystemInterface(name="electrics", mass_kg=fresh.mass_kg * 1.004, voltage_v=fresh.voltage_v * 1.006,
                               peak_current_a=None, power_draw_w=fresh.power_draw_w + 2e-7, cg_x_mm=1.0,
                               cg_y_mm=2.0, cg_z_mm=3.0, env_x_mm=4.0, env_y_mm=5.0, env_z_mm=6.0,
                               max_temp_c=60.0, cooling_airflow_cms=0.2, heat_reject_w=99.0,
                               env_origin_mm=(1.0, 2.0, 3.0), mount_load_n=500.0, mount_points=4,
                               mounts_on="chassis"))
    res["drift"] = [pdw.sync_ledger(led, pack), led.interfaces["electrics"]]
    led = IntegrationLedger()
    led.set(SubsystemInterface(name="hv", mass_kg=1.0, voltage_v=1.0, peak_current_a=1.0, power_draw_w=1.0,
                               heat_reject_w=5.0, cg_x_mm=7.0))
    res["named"] = [pdw.sync_ledger(led, pdw.PackSpec(), name="hv", heat_reject_w=8.0), led.interfaces["hv"]]
    led = IntegrationLedger()
    led.set(SubsystemInterface(name="electrics", mass_kg=0.0, voltage_v=fresh.voltage_v, peak_current_a=0.0))
    res["zero"] = [pdw.sync_ledger(led, pdw.PackSpec(n_series=0)), led.interfaces["electrics"]]
    check("power_draw_ledger", res)


def test_audit_end_to_end(tmp_path):
    bad = _cache(_save(tmp_path, "bad", {**_variant(), "Extra": {}}, pack=PACK_ROWS + [
        ("Internal Resistance Battery Pack (Ohms)", "=1.792+0"), ("Pack Mass (kg)", 30.0),
        ("Joule Heating Max (kWh)", 4.48)], first="Cover"), {"1.792+0": 1.792, "1/8": 0.125})
    ok_r = pdw.read_pack_config(_save(tmp_path, "probe_r", {}, pack=PACK_ROWS)).resistance_ohm()
    clean = _save(tmp_path, "clean", {}, pack=PACK_ROWS + [
        ("Pack Internal Resistance (ohm)", ok_r), ("Joule heating (W)", 1.0)], first="Cover")
    nopack = _save(tmp_path, "nopack", {"Notes": {"A1": "hello"}, "More": {"A1": "x"}})
    res = {}
    for name, p, kw in [("bad", bad, {}), ("clean", clean, {"drive": pdw.DriveSpec(motor_peak_power_kw=10.0)}),
                        ("nopack", nopack, {}), ("bad_nr", bad, {"include_readers": False})]:
        a = pdw.audit(p, **kw)
        res[name] = {"a": a, "md": a.to_markdown(), "b": a.blocking(), "w": a.warnings()}
    res["kin_nopack"] = pdw.audit_kinematik_readers(nopack)
    check("power_draw_audit", _norm(res, tmp_path))


def test_kinematik_reader_fallbacks(tmp_path, monkeypatch):
    import sys
    import types
    from suspension import ev_excel_roundtrip as rt
    p = _save(tmp_path, "kin", {}, pack=PACK_ROWS)
    flat = types.SimpleNamespace(current_draw_a=[5.0] * 40)
    monkeypatch.setattr(rt, "lap_to_excel_roundtrip_from_db", lambda v, t, params: flat)
    res = {"mirror": pdw.audit_kinematik_readers(p)}
    monkeypatch.setitem(sys.modules, "suspension.ev_electrical_check", None)
    res["no_elec"] = pdw.audit_kinematik_readers(p)
    check("power_draw_kin", _norm(res, tmp_path))


def _synthetic(n=50, dt=1.0, i_pk=50.0, p=1000.0, joule=10.0, tq=50.0, rpm=1000.0, disc=0, clamped=0, sw=1,
               infeasible=0, pack=None, drive=None):
    pack = pack or pdw.PackSpec()
    drive = drive or pdw.DriveSpec()
    d = [(k, 30.0, 2.5 + k) for k in range(disc)]
    return pdw.PowerDrawTrace(dt, [10.0] * n, [0.0] * n, [1.0] * n, [1.0] * n, [1.0] * n, [p] * n, [p] * n,
                              [i_pk] + [0.0] * (n - 1), [joule] * n, [rpm] * n, [tq] * n, infeasible, clamped,
                              d, sw, pack, pdw.VehicleSpec(), drive)


def test_trace_findings_boundaries():
    """Hand-built traces sitting exactly on each threshold, so every comparison's direction is pinned."""
    one = pdw.PackSpec(n_series=100, n_parallel=1, cell_voltage_v=4.0, cell_capacity_ah=2.5)   # 1.0 kWh
    usable = one.energy_kwh() * 0.8 - one.energy_kwh() * 0.1
    p_at = usable * 3.6e6
    cases = {"fuse_eq": dict(i_pk=50.0), "fuse_over": dict(i_pk=50.0000001), "fuse_under": dict(i_pk=49.0),
             "energy_mid": dict(n=1, p=0.8 * 3.6e6, pack=one),
             "energy_over": dict(n=1, p=p_at * 1.0001, pack=one), "tq_peak": dict(tq=120.0),
             "tq_cont": dict(tq=77.0), "tq_mid": dict(tq=100.0), "tq_over": dict(tq=121.0),
             "rpm_eq": dict(rpm=23500.0), "rpm_over": dict(rpm=23501.0),
             "disc_2pct": dict(disc=1), "disc_more": dict(disc=2), "disc_one_of_10": dict(n=10, disc=1),
             "clamped": dict(clamped=3), "clamped_s2": dict(clamped=3, sw=2), "clamped_s3": dict(clamped=4, sw=3),
             "infeasible": dict(infeasible=5), "big": dict(dt=2.0, joule=1e6, disc=1, i_pk=123.4)}
    res = {k: pdw.trace_findings(_synthetic(**kw)) for k, kw in cases.items()}
    check("power_draw_bounds", res)


def test_markdown_ordering_and_escaping():
    S = pdw.Severity
    F = pdw.Finding
    fs = [F("ok", S.OK, "fine"), F("info", S.INFO, "a | b"), F("miss", S.MISSING, "line1\nline2"),
          F("fail", S.FAIL, "x|y\nz"), F("warn", S.WARN, "w"), F("fail2", S.FAIL, "f2"), F("info2", S.INFO, "i2"),
          F("ok2", S.OK, "o2"), F("miss2", S.MISSING, "m2"), F("warn2", S.WARN, "w2")]
    packs = [pdw.PackSpec(), pdw.PackSpec(n_series=96, n_parallel=2, fuse_max_a=150.0, cell_resistance_ohm=0.02),
             pdw.PackSpec(n_series=150, n_parallel=5, cell_voltage_v=3.65, fuse_max_a=333.0),
             pdw.PackSpec(n_series=72, n_parallel=6, fuse_max_a=512.0, cell_resistance_ohm=0.003)]
    packs.append(pdw.PackSpec(n_series=600, n_parallel=8, fuse_max_a=900.0, cell_resistance_ohm=0.005))
    assert any(f"{p.power_at_fuse_limit_w() / 1000:.1f}" != f"{p.power_at_fuse_limit_w() / 1001:.1f}" for p in packs)
    assert any(f"{p.max_deliverable_power_w() / 1000:.0f}" != f"{p.max_deliverable_power_w() / 1001:.0f}"
               for p in packs)
    res = {f"p{i}": pdw.PowerDrawAudit("x.xlsx", p, fs).to_markdown() for i, p in enumerate(packs)}
    res["nopack"] = pdw.PowerDrawAudit("x.xlsx", None, fs[::-1]).to_markdown()
    res["order"] = pdw._SEVERITY_ORDER
    check("power_draw_md", res)
