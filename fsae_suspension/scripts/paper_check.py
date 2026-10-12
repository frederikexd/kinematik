# ============================================================================
#  KinematiK — Formula SAE suspension & vehicle dynamics toolkit
#  Created by Frederik Thio. Copyright (c) 2026 Frederik Thio.
#  Open source. Original author: Frederik Thio, creator of KinematiK.
# ============================================================================
"""Claim-by-claim check of the paper and its supplement against the published manifests.

Each claim names where it is made, the text the document prints (which must carry the number) and the code that
recomputes it from the manifests. The check passes a claim when the recomputed value, rounded as printed, appears in
the quoted text; with the documents given it also checks that the quoted text is in them, so a number edited in one
place and not the other fails.

usage: python scripts/paper_check.py [paper.docx|.txt supplement.docx|.txt] [--json out.json]
Lengths mm, forces N, angles deg.
"""
from __future__ import annotations

import json
import math
import re
import sys
from dataclasses import dataclass
from functools import cached_property
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from suspension import (actuation as act, dynamic_ackermann as da, elastokinematics as ek,      # noqa: E402
                        genesis_analysis as ga, inverse_genesis as ig, robustness as rb,
                        target_derivation as td, tire_field as tf)
from suspension.genesis_repro import GenesisManifest, hp_from_dict, vehicle_load_cases     # noqa: E402
from suspension.ghost_topology import _rc_height_mm                                        # noqa: E402
from suspension.kinematics import SuspensionKinematics                                     # noqa: E402

MAN = ROOT / "manifests"
DAMPER_EXTENDED_MM = 200.0      # Ohlins TTX25 MkII FSAE, 200 mm eye to eye extended, 57 mm stroke (scripts/final_corners_v3.DAMPER)
FRONT, REAR = "front_v11_final", "rear_v11_final"
TRACK = 1210.0
REAR_KW = {"traction_g": 1.06, "inboard_drive": True, "inboard_brakes": True}


def manifest(name: str) -> GenesisManifest:
    return GenesisManifest.from_json((MAN / f"{name}.recorded.genesis.json").read_text())


def winner(name: str):
    return hp_from_dict(manifest(name).recorded["winner_hardpoints"])


class Ctx:
    """Every recomputed quantity the claims draw on, computed once."""

    def __init__(self, front: str = FRONT, rear: str = REAR):
        self.front, self.rear = front, rear
        self.mF, self.mR = manifest(front), manifest(rear)
        self.hF, self.hR = winner(front), winner(rear)
        self.vF, self.vR = self.mF.objects()[2], self.mR.objects()[2]

    # ---- walls and kinematics ------------------------------------------------ #
    @cached_property
    def walls(self) -> dict:
        return {"F": self.vF.properties.evaluate(self.hF), "R": self.vR.properties.evaluate(self.hR),
                "F_viol": self.vF.wall_violations(self.hF), "R_viol": self.vR.wall_violations(self.hR),
                "F_env": self.vF.envelope_margin(self.hF), "R_env": self.vR.envelope_margin(self.hR)}

    @staticmethod
    def _corner(hp) -> dict:
        tr = np.linspace(-25, 25, 11)
        k = SuspensionKinematics(hp)
        sts = [k.solve_at_travel(float(t)) for t in tr]
        rc = np.array([_rc_height_mm(s, TRACK) for s in sts])
        cam = np.array([s.camber for s in sts]); toe = np.array([s.toe for s in sts])
        return {"rc": float(rc[5]), "mig": float(np.polyfit(tr, rc, 1)[0]), "gain": float(np.polyfit(tr, cam, 1)[0]),
                "bump_steer": float(toe.max() - toe.min())}

    @staticmethod
    def _anti(hp, prop, **kw) -> float:
        b = ig.SolvedPropertyBounds(bounds=[ig.PropertyBound(prop, lo=-1e9)], **kw)
        return ig.properties_of(hp, b, only=[prop])[prop]

    @cached_property
    def corners(self) -> dict:
        F, R = self._corner(self.hF), self._corner(self.hR)
        F["ad73"] = self._anti(self.hF, "anti_dive_pct", brake_bias_front=0.73)
        F["ad60"] = self._anti(self.hF, "anti_dive_pct", brake_bias_front=0.60)
        R["as"] = self._anti(self.hR, "anti_squat_pct")
        R["al_in"] = self._anti(self.hR, "anti_lift_pct", brake_bias_front=0.73, rear_brakes_inboard=True)
        R["al_out"] = self._anti(self.hR, "anti_lift_pct", brake_bias_front=0.73, rear_brakes_inboard=False)
        F["share"] = td.front_lltd_share(F["rc"], R["rc"], 0.522)
        return {"F": F, "R": R}

    @cached_property
    def targets(self) -> dict:
        F, R = self.corners["F"], self.corners["R"]
        v = td.Vehicle(); t = td.Tire(); a = td.design_accelerations(t)
        from dataclasses import replace
        return {"f_brake": td.f_min(v, F["ad73"], a["straight"], 0.0, "front"),
                "f_comb_F": td.f_min(v, F["ad73"], a["combined"], a["combined"], "front"),
                "f_trac": td.f_min(v, R["as"], a["straight"], 0.0, "rear"),
                "f_comb_R": td.f_min(v, R["as"], a["combined"], a["combined"], "rear"),
                "ride_F": td.ride_frequency_derived(v, t, F["ad73"], "front"),
                "ride_R": td.ride_frequency_derived(v, t, R["as"], "rear"),
                "mu_F": tf.mu_threshold_for("anti_dive_floor_pct", F["ad73"], veh=replace(v, ride_front_hz=2.6)),
                "mu_R": tf.mu_threshold_for("anti_squat_floor_pct", R["as"])}

    @cached_property
    def pitch(self) -> dict:
        F, R = self.corners["F"], self.corners["R"]

        def p(ad, al, a_g):
            dF = 300 * 9.81 * a_g * 0.280 / (2 * 1.630)
            kf = 60.1 * (2 * math.pi * 2.6) ** 2 / 1000; kr = 66.1 * (2 * math.pi * 3.0) ** 2 / 1000
            dive = dF * (1 - ad / 100) / kf; rise = dF * (1 - al / 100) / kr
            return dive, rise, math.degrees(math.atan((dive + rise) / 1630)), dF * (1 - al / 100)
        dive, rise, deg, unload = p(F["ad73"], R["al_in"], 1.5)
        s0 = F["share"]
        body = td.front_lltd_share(F["rc"] + F["mig"] * dive, R["rc"] - R["mig"] * rise, 0.522)
        return {"dive": dive, "rise": rise, "deg": deg, "unload": unload, "deg18": p(F["ad73"], R["al_in"], 1.8)[2],
                "shift_body_pts": 100 * (body - s0)}

    # ---- steering -------------------------------------------------------------- #
    @cached_property
    def steering(self) -> dict:
        hp, se = self.hF, self.vF.properties.steer_effort
        rc = se.kinematic_ratio(hp); sw = se.lock_sw_deg(hp)
        r = da.lock_report(hp, se.rack_c_mm_per_rev, se.lock_road_wheel_deg, lock_basis="mean")
        rack = sw * se.rack_c_mm_per_rev / 360.0
        jm = ig.joint_misalignment(hp, rack_travel_mm=rack)
        return {"ratio_centre": rc, "lock_sw": sw, "ratio_avg": sw / se.lock_road_wheel_deg, "inner": r["inner_deg"],
                "outer": r["outer_deg"], "ack": r["ackermann_pct_lo_and_lock"][1], "ack_lo": r["ackermann_pct_lo_and_lock"][0],
                "cam_out": r["camber_at_lock_outside_deg"], "cam_in": r["camber_at_lock_inside_deg"], "rack": rack,
                "swing": ig.joint_swing(hp, rack_travel_mm=rack)["_worst"][1], "misalign": jm["_worst"][1],
                "effort_margin": self.walls["F"].get("steer_effort_margin")}

    # ---- actuation ----------------------------------------------------------------- #
    @cached_property
    def actuation(self) -> dict:
        out = {}
        for k, hp in (("F", self.hF), ("R", self.hR)):
            m = act.actuation_metrics(hp)
            ot = rb.overtravel_check(hp, damper_extended_mm=DAMPER_EXTENDED_MM)
            ax = np.asarray(hp.rocker_axis, float)
            out[k] = {"mr": m.mr_static, "spread": m.spread, "tr": tuple(m.transmission_deg), "oop": m.out_of_plane_deg,
                      "stroke": m.stroke_used_mm, "tilt": math.degrees(math.acos(min(1.0, abs(ax[0]) / np.linalg.norm(ax)))),
                      "bracket_mm": float(np.linalg.norm(np.asarray(hp.pushrod_outer, float) - np.asarray(hp.lower_outer, float))),
                      "ot_bottom": ot.get("damper_bottoms_at_mm"), "ot_tr": (ot.get("transmission_min_deg"), ot.get("transmission_max_deg")),
                      "ot_mr": (ot.get("mr_max"), ot.get("mr_min")), "ot_stroke": ot.get("stroke_used_mm"),
                      "ot_avail": ot.get("stroke_available_mm"), "spring_mm": tuple(m.spring_length_mm)}
        return out

    # ---- structure ------------------------------------------------------------------ #
    @cached_property
    def structure(self) -> dict:
        from suspension import genesis_walls as gw
        sF = self.vF.properties.structure; sR = self.vR.properties.structure
        fF, fR = sF.margins(self.hF), sR.margins(self.hR)
        bare = gw.LinkStructure(cases=sF.cases).margins(self.hF)
        return {"F": fF, "R": fR, "F_bare": bare}

    @staticmethod
    def first_yield(P: float, L: float, od: float, wall: float, sy: float = 435.0, e: float = 205000.0,
                    crook: float = 1.0 / 1000.0) -> float:
        """Load factor to first yield of a crooked strut in compression: the factor on P at which the Perry-Robertson
        peak stress, with the crookedness bending growing as the load nears P_cr, reaches sy (dimensionless)."""
        from scipy.optimize import brentq
        A = math.pi / 4 * (od ** 2 - (od - 2 * wall) ** 2); I = math.pi / 64 * (od ** 4 - (od - 2 * wall) ** 4)
        pcr = math.pi ** 2 * e * I / L ** 2
        f = lambda lam: lam * P / A * (1 + crook * L * (od / 2) / (I / A) / (1 - lam * P / pcr)) - sy   # noqa: E731
        return brentq(f, 1e-6, pcr / P * (1 - 1e-9))

    @cached_property
    def load_factors(self) -> dict:
        """Load factor to first yield of the compressed legs Table 8 names (front lower-rear on its section and on
        5/8 x 0.035 in; rear lower-front and upper-rear)."""
        s = self.structure
        fl = s["F"]["LR"]; rl = s["R"]["LF"]; ru = s["R"]["UR"]
        return {"F_LR": self.first_yield(-fl["worst_n"], fl["length_mm"], fl["od_mm"], fl["wall_mm"]),
                "F_LR_bare": self.first_yield(-fl["worst_n"], fl["length_mm"], 15.875, 0.889),
                "R_LF": self.first_yield(-rl["worst_n"], rl["length_mm"], rl["od_mm"], rl["wall_mm"]),
                "R_UR": self.first_yield(-ru["worst_n"], ru["length_mm"], ru["od_mm"], ru["wall_mm"])}

    @cached_property
    def pushrod(self) -> dict:
        """Pushrod (5/8 x 0.035 in) worst compression over the shock cases, buckling and Perry-Robertson factors."""
        from suspension import genesis_walls as gw
        from suspension.loadpath import solve_member_forces
        out = {}
        for k, hp, vol in (("F", self.hF, self.vF), ("R", self.hR, self.vR)):
            st = vol.properties.structure
            kin = SuspensionKinematics(hp); s0 = kin.solve_at_travel(0.0)
            P = max(-solve_member_forces(kin, s0, gw._wheel_load(v)).forces["PR"] for v in st.cases.values())
            L = float(np.linalg.norm(np.asarray(hp.pushrod_outer, float) - np.asarray(hp.rocker_pushrod, float)))
            od, w = 15.875, 0.889
            A = math.pi / 4 * (od ** 2 - (od - 2 * w) ** 2); I = math.pi / 64 * (od ** 4 - (od - 2 * w) ** 4)
            pcr = math.pi ** 2 * st.e_mpa * I / L ** 2
            smax = P / A * (1 + st.crook * L * (od / 2) / (I / A) / (1 - P / pcr))
            out[k] = {"P": P, "L": L, "buckling": pcr / P, "yield": st.sy_mpa / smax}
        return out

    @cached_property
    def bracket(self) -> dict:
        """The front lower-fore bracket: the bound (whole leg load across both plates) and the load resolved along and
        across the bolt (paper section 5.5), at 13 mm reach on 50 x 6.35 mm plates."""
        hp = self.hF
        F = abs(self.structure["F"]["LF"]["worst_n"])
        A, B, O = (np.asarray(getattr(hp, p), float) for p in ("lower_front_inner", "lower_rear_inner", "lower_outer"))
        c = abs(float(((O - A) / np.linalg.norm(O - A)) @ ((B - A) / np.linalg.norm(B - A))))
        fa, fp = F * c, F * math.sqrt(1 - c * c)
        r, w, t = 13.0, 50.0, 6.35
        s = 1.28 * (fa / 2 * r) / (w * t * t / 6) + 1.28 * (fp / 2 * r) / (t * w * w / 6) + 1.28 * (fp / 2) / (w * t)
        return {"F": F, "bound": ga.bracket_fos(F, r, w, t)["fos"], "required_kN": 2.0 * F / 1000.0,
                "angle": math.degrees(math.acos(c)), "along_kN": fa / 1000, "inplane_kN": fp / 1000, "resolved": 460.0 / s}

    @cached_property
    def bearings(self) -> dict:
        out = {}
        for k, hp, axle, kw in (("F", self.hF, "front", {}), ("R", self.hR, "rear", REAR_KW)):
            cases = list(vehicle_load_cases(axle=axle, brake_bias_front=0.73, **kw))
            bw = (self.vF if k == "F" else self.vR).properties.bearing
            req = act.required_rating_vs_spacing(hp, cases, [30.0, 60.0])
            out[k] = {"s0_6001": min(r["s0"] for r in act.rocker_bearing_loads(hp, cases)), "s0": bw.s0_min(hp),
                      "C0_30": req[0]["required_C0_N"], "C0_60": req[1]["required_C0_N"]}
        return out

    # ---- compliance ------------------------------------------------------------------ #
    def _members(self, vol):
        el = vol.properties.elasto
        return dict(el.members) if el is not None and el.members else None

    @cached_property
    def compliance(self) -> dict:
        """Front: own-corner derived flexibility (the manifest's compliance wall), links alone and 30 kN/mm, every case."""
        from suspension.compliance import MemberStiffness
        el = self.vF.properties.elasto
        mem = dict(el.members)
        out = {}
        for c in vehicle_load_cases(axle="front", brake_bias_front=0.73):
            kw = dict(Fx=c.Fx, Fy=c.Fy, Fz=c.Fz, Mz=c.mz_Nmm, Fx_wc=getattr(c, "Fx_wc", 0.0))
            out[c.name] = {
                "derived": ek.solve_elastokinematic(self.hF, ek.ElastoSpec(stiffness=el.stiffness, members=mem, **kw),
                                                    jacobian=False).change,
                "links": ek.solve_elastokinematic(self.hF, ek.ElastoSpec(stiffness=ek.StiffnessField(1e9), members=mem, **kw),
                                                  jacobian=False).change,
                "k30": ek.solve_elastokinematic(self.hF, ek.ElastoSpec(stiffness=ek.StiffnessField(30000.0), members=mem, **kw),
                                                jacobian=False).change}
        return out

    @cached_property
    def frame(self) -> dict:
        """scripts/frame_derivation.py on the final corners (about 90 s): torsion, pickup flexibility, whole car."""
        import frame_derivation as fd
        fd.MAN = MAN
        fd.CORNERS = {"front": self.front, "rear": self.rear}
        return fd.run()

    @cached_property
    def compliance_table(self) -> dict:
        """Table S6.36: outer-wheel toe and camber per case at 30 kN/mm, links alone, the derived own-corner
        flexibility and the whole car (frame part from the free-frame solve plus the links' own give)."""
        from suspension.compliance import MemberStiffness
        fr = self.frame
        walls = {"F": self.vF.properties.elasto.members,
                 "R": {k: MemberStiffness(material="Steel 4130", od_mm=15.875,
                                          wall_mm={"LF": 1.651, "LR": 1.245}.get(k, 0.889)) for k in ("UF", "UR", "LF", "LR", "TR")}}
        fv = {"1.5g corner": "1.5g corner", "1.5g braking": "1.5g braking", "combined 1.06g": "combined 1.06g",
              "combined 1.06g traction": "combined 1.06g traction"}
        out = {}
        for k, hp, axle, kw in (("F", self.hF, "front", {}), ("R", self.hR, "rear", REAR_KW)):
            C = np.asarray(fr["corners"][axle]["flexibility_mm_per_N"], float)
            stiff = {"k30": ek.StiffnessField(30000.0), "links": ek.StiffnessField(1e9),
                     "derived": ek.StiffnessField(30000.0, coupled=C.tolist())}
            for c in vehicle_load_cases(axle=axle, brake_bias_front=0.73, **kw):
                lk = dict(Fx=c.Fx, Fy=c.Fy, Fz=c.Fz, Mz=c.mz_Nmm, Fx_wc=getattr(c, "Fx_wc", 0.0))
                row = {n: ek.solve_elastokinematic(hp, ek.ElastoSpec(stiffness=s, members=dict(walls[k]), **lk),
                                                   jacobian=False).change for n, s in stiff.items()}
                if c.name in fv:
                    ch = fr["full_vehicle"][fv[c.name]]["corners"][k + "R"]["channels"]
                    row["whole"] = {q: ch[q] + row["links"][q] for q in ("toe", "camber")}
                    row["frame_part"] = dict(ch)
                out[(k, c.name)] = row
        return out

    @cached_property
    def packaging(self) -> dict:
        """Table 7 and section 5.5 extras: chassis mounts to the nearest ChassisV10 node, swept clearance to the tubes
        alone and to tubes plus envelopes, bearing-hub clearance to the tubes, spring rate and ride frequency across
        travel, spring roll-stiffness change at 12.4 mm of bump, per-corner tube mass and, at the front, the spring
        length change with the rack at 27 mm each way."""
        import actuation_package as ap
        from suspension.genesis_analysis import frame_graph_from_step, parse_step_tubes
        if not ap.STEP.exists():
            raise FileNotFoundError(f"ChassisV10.STEP not found ({ap.STEP}); set KINEMATIK_V10_STEP")
        from suspension.genesis_repro import nodes_from_framegraph
        fg = frame_graph_from_step(parse_step_tubes(ap.STEP.read_text()))
        out = {}
        hubs = {"F": (47.0 / 2 + 3.0, 22.0), "R": (52.0 / 2 + 3.5, 22.0)}       # hub radius (6204 + 3 mm wall, 6205 + 3.5 mm); half of 44 mm
        for k, hp, axle in (("F", self.hF, "front"), ("R", self.hR, "rear")):
            nodes = np.asarray(nodes_from_framegraph(fg, ap.AZ[axle], ap.G), float)
            reach = max(float(np.min(np.linalg.norm(nodes - np.asarray(getattr(hp, p), float), axis=1)))
                        for p in ("rocker_pivot", "spring_inner"))
            frame = ap.frame_obstacles(axle)
            wheel, chassis = ap.envelopes(hp, axle)
            kw = dict(pushrod_skip_mm=ap.SKIP, spring_skip_mm=ap.EYE, member_radii=(ap.R_PR, ap.R_ROCK, ap.R_ROCK, ap.R_COIL))
            m_t = act.actuation_metrics(hp, obstacles=frame, **kw)
            m_a = act.actuation_metrics(hp, obstacles=frame + chassis, wheel_obstacles=wheel, **kw)
            ax = np.asarray(hp.rocker_axis, float); ax /= np.linalg.norm(ax)
            c = np.asarray(hp.rocker_pivot, float); rh, hl = hubs[k]
            hub = min(act._seg_dist(c - hl * ax, c + hl * ax, a, b) - rh - r for a, b, r in frame)
            ms, hz = ap.SPRUNG[axle]
            k_w = (2 * math.pi * hz) ** 2 * ms / 1000.0
            m0 = act.actuation_metrics(hp)
            mr = np.asarray(m0.mr, float); trv = np.asarray(m0.travel, float)
            mr_b = float(np.interp(12.4, trv, mr))
            L = {p: float(np.linalg.norm(np.asarray(getattr(hp, a), float) - np.asarray(getattr(hp, b), float)))
                 for p, (a, b) in {"UF": ("upper_outer", "upper_front_inner"), "UR": ("upper_outer", "upper_rear_inner"),
                                   "LF": ("lower_outer", "lower_front_inner"), "LR": ("lower_outer", "lower_rear_inner"),
                                   "TR": ("tie_rod_outer", "tie_rod_inner"), "PR": ("pushrod_outer", "rocker_pushrod")}.items()}
            vol = self.vF if k == "F" else self.vR
            pm = dict(vol.properties.structure.per_member)

            def mass(sec):
                g = 0.0
                for p, l in L.items():
                    od, w = sec.get(p, (15.875, 0.889))
                    g += 7.85e-3 * math.pi / 4 * (od ** 2 - (od - 2 * w) ** 2) * l
                return g
            bare = {p: v for p, v in pm.items() if p == "TR"}
            out[k] = {"node_reach": reach, "clear_tubes": m_t.clearance_mm, "clear_all": m_a.clearance_mm, "hub": hub,
                      "spring": k_w / m0.mr_static ** 2, "ride": (float(hz * mr.min() / m0.mr_static), float(hz * mr.max() / m0.mr_static)),
                      "roll_change_pct": 100 * ((mr_b / m0.mr_static) ** 2 - 1), "mass_g": mass(pm), "mass_bare_g": mass(bare),
                      "lengths": L,
                      "steer_spring": ap.steer_coupling(hp, 27.0) if k == "F" else None}
        return out

    @cached_property
    def earlier(self) -> dict:
        """The delivered front and the released ChassisV11 fronts (pushrod on the lower wishbone) on the corrected statics:
        lower-leg factors on 5/8 x 0.035 in under the shock cases, the delivered lower-fore bracket, and the refusal starts."""
        from suspension import genesis_walls as gw
        d = winner("front_steering_walls")
        m73 = gw.LinkStructure.from_shock_cases("front", 0.73).margins(d)
        m60 = gw.LinkStructure.from_shock_cases("front", 0.60).margins(d)
        out = {"d73": m73, "d60": m60}
        for n in ("front_v11_29deg", "front_v11_25deg"):
            out[n] = gw.LinkStructure.from_shock_cases("front", 0.73).margins(winner(n))
        F = abs(m73["LF"]["worst_n"])
        out["bracket_built"] = ga.bracket_fos(F, 27.9, 30, 5)["fos"]; out["bracket_node"] = ga.bracket_fos(F, 13.0, 50, 6.35)["fos"]
        out["F_LF"] = F
        for n in ("front_v11_29deg_b73_tube035", "front_v11_29deg_b73_6001"):
            mm = manifest(n); vol = mm.objects()[2]; h = hp_from_dict(mm.hardpoints)
            out[n] = vol.properties.evaluate(h)
            if n.endswith("tube035"):
                out["tube035_LR"] = vol.properties.structure.margins(h)["LR"]
        return out

    @cached_property
    def damper_front(self) -> dict:
        k = SuspensionKinematics(self.hF)
        L = {t: k.spring_length_at(k.solve_at_travel(float(t)), seed=0.0)[0] for t in (-25.0, 0.0, 25.0)}
        return {"eye": L[0.0], "extended_min": L[-25.0], "extended_max": L[25.0] + 57.0}

    @cached_property
    def baseline(self) -> dict:
        """front_v11_curves_only's in-band candidates judged against front_v11_final's walls (method_benchmarks)."""
        import method_benchmarks as mb
        mb.MAN = MAN
        hits = [c for c in mb.judge_curves_only()["candidates"] if c["hit"]]
        win = next(c for c in hits if c["winner"])
        others = [c for c in hits if not c["winner"]]
        return {"winner": win["properties"], "others_link": [c["properties"]["link_fos_min"] for c in others],
                "others_ad": [c["properties"]["anti_dive_pct"] for c in others]}

    @cached_property
    def lf_share(self) -> float:
        """Toe in the 1.5 g stop from the lower-front leg's own stretch alone (every other member and pickup rigid)."""
        el = self.vF.properties.elasto
        c = [c for c in vehicle_load_cases(axle="front", brake_bias_front=0.73) if c.name == "1.5g braking"][0]
        mem = {"LF": el.members["LF"]}
        return ek.solve_elastokinematic(self.hF, ek.ElastoSpec(stiffness=ek.StiffnessField(1e9), members=mem, Fx=c.Fx, Fy=c.Fy,
                                                                Fz=c.Fz, Mz=c.mz_Nmm, Fx_wc=getattr(c, "Fx_wc", 0.0)),
                                        jacobian=False).change["toe"]

    @cached_property
    def corner_load(self) -> float:
        """Outer front load at 1.5 g with 442 N downforce, +/-2 deg bank and grade and steer warp, the adverse sign."""
        w = self.warp["front_pair_N"]
        return max(td.corner_loads(td.Vehicle(), a_lat_g=1.5, bank_deg=2.0, grade_deg=-2.0, downforce_n=442.0,
                                   steer_warp_n=s * w)["total"]["FR"] for s in (1.0, -1.0))

    @cached_property
    def warp(self) -> dict:
        w = td.steer_warp(self.hF, self.steering["rack"], 454.0, 405.0, TRACK, TRACK)
        return w

    @cached_property
    def twist(self) -> dict:
        J = ek.compliance_jacobian(self.hF, points=("tie_rod_inner",))
        u = np.asarray(self.hF.tie_rod_outer, float) - np.asarray(self.hF.tie_rod_inner, float); u /= np.linalg.norm(u)
        tpm = float(abs(J[1] @ u))
        return {"toe_per_mm": tpm, "one_flat": 0.176 * tpm,
                "t618": ek.frame_twist_toe(618.0, 2000.0, 1630.0, 250.0, 200.0, tpm)}


@dataclass
class Claim:
    id: str
    where: str            # "P §5.2" or "S S6.37"
    quote: str            # text as printed; must contain the formatted value
    value: object         # callable(ctx) -> float
    fmt: str = "{:.1f}"
    doc: str = "P"        # P paper, S supplement, "PS" both

    def evaluate(self, ctx: Ctx) -> dict:
        try:
            v = self.value(ctx)
            s = self.fmt.format(v)
            s_alt = s.replace("-", "−")
            ok = (s in self.quote) or (s_alt in self.quote)
            if not ok and s.startswith("-"):                  # printed as a magnitude where the sign is in the words
                ok = s[1:] in self.quote
            return {"id": self.id, "where": self.where, "value": v, "printed": s, "ok": ok, "quote": self.quote}
        except FileNotFoundError as e:                          # an input outside the repository: not judged
            return {"id": self.id, "where": self.where, "value": None, "printed": "skipped", "ok": None,
                    "quote": self.quote, "note": str(e)}
        except Exception as e:                                  # a claim whose code fails is a failed claim
            return {"id": self.id, "where": self.where, "value": None, "printed": f"error: {e}", "ok": False,
                    "quote": self.quote}


def _norm(t: str) -> str:
    t = t.replace("−", "-").replace("–", "-").replace(" ", " ").replace(" ", " ")
    return re.sub(r"\s+", " ", t)


def doc_text(path: str | None) -> str | None:
    if not path:
        return None
    p = Path(path)
    if p.suffix == ".docx":
        from docx import Document
        d = Document(str(p))
        parts = [q.text for q in d.paragraphs]
        for tb in d.tables:
            for row in tb.rows:
                for c in row.cells:
                    parts.append(c.text)
        return _norm("\n".join(parts))
    return _norm(p.read_text())


def run(claims, paper: str | None = None, supp: str | None = None, ctx: Ctx | None = None) -> list[dict]:
    ctx = ctx or Ctx()
    texts = {"P": doc_text(paper), "S": doc_text(supp)}
    rows = []
    for c in claims:
        r = c.evaluate(ctx)
        for d in c.doc:
            t = texts.get(d)
            if t is not None:
                found = _norm(c.quote) in t
                r[f"in_{d}"] = found
                if r["ok"] is not None:
                    r["ok"] = r["ok"] and found
        rows.append(r)
    return rows


def report(rows) -> str:
    w = max(len(r["id"]) for r in rows)
    lines = [f"{'claim':{w}}  {'where':16} {'printed':>10}  result   quote"]
    for r in rows:
        res = "SKIP" if r["ok"] is None else ("PASS" if r["ok"] else "FAIL")
        lines.append(f"{r['id']:{w}}  {r['where']:16} {r['printed']:>10}  {res:6}   {r['quote'][:110]}")
    n = sum(r["ok"] is True for r in rows); k = sum(r["ok"] is None for r in rows)
    lines.append(f"{n} of {len(rows)} claims pass" + (f", {k} skipped (ChassisV10.STEP not found)" if k else ""))
    return "\n".join(lines)


if __name__ == "__main__":
    from paper_claims import CLAIMS                     # the claim list lives beside this script
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    out = sys.argv[sys.argv.index("--json") + 1] if "--json" in sys.argv else None
    if out in args:
        args.remove(out)
    rows = run(CLAIMS, *(args + [None, None])[:2])
    print(report(rows))
    if out:
        Path(out).write_text(json.dumps(rows, indent=1, default=str))
    sys.exit(0 if all(r["ok"] is not False for r in rows) else 1)
